"""Behavioral coverage for the compact stage-14 editor shell."""

from pathlib import Path
import shutil
import pytest

from PySide6.QtCore import QEvent, QPoint, QSettings, QSize, Qt, QTimer
from PySide6.QtGui import QHelpEvent, QIcon, QImage, QPixmap
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QDialog, QFileDialog, QMessageBox, QTreeView

from componentpress.bootstrap import project_service
from componentpress.editor.component_dialog import ComponentDialog
from componentpress.editor.controllers import ProjectController
from componentpress.editor.main_window import MainWindow
from componentpress.editor.ui_state import UiState
from componentpress.editor.qt_models import ID_ROLE
from componentpress.editor.project_settings_dialog import ProjectSettingsDialog
from componentpress.editor import main_window as main_window_module
from componentpress.editor.dialogs import format_error
from componentpress.project_io.repository import FileProjectRepository


DEMO = Path(__file__).resolve().parents[1] / "examples" / "demo-game"


def make_window(tmp_path: Path, qtbot):
    QSettings.setDefaultFormat(QSettings.Format.IniFormat)
    QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, str(tmp_path / "qsettings"))
    root = tmp_path / "game"
    shutil.copytree(DEMO, root)
    service = project_service()
    service.register(root, "second", "Второй компонент")
    window = MainWindow(ProjectController(service))
    window.ui_state = UiState(QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat))
    window.load_session(service.open_session(root))
    qtbot.addWidget(window)
    window.show()
    return root, window, service


def use_temp_ui_settings(window: MainWindow, tmp_path: Path) -> None:
    window.ui_state = UiState(QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat))


@pytest.mark.parametrize(("locale", "help_title"), (("ru", "Справка"), ("en", "Help")))
def test_startup_menus_receive_clicks_without_project_and_after_restore(
    tmp_path: Path, qtbot, monkeypatch, locale: str, help_title: str
) -> None:
    from componentpress.i18n import language, set_language

    previous_language = language()
    settings_path = tmp_path / "settings.ini"
    monkeypatch.setattr(
        main_window_module,
        "UiState",
        lambda: UiState(QSettings(str(settings_path), QSettings.Format.IniFormat)),
    )
    set_language(locale)
    try:
        window = MainWindow(ProjectController(project_service()))
        qtbot.addWidget(window)
        window.show()
        qtbot.waitUntil(window.menuBar().isVisible)
        bar = window.menuBar()
        expected_menus = (
            ["Файл", "Правка", "Проект", "Вид", "Сборка", "Настройки", "Справка"]
            if locale == "ru"
            else ["File", "Edit", "Project", "View", "Build", "Settings", "Help"]
        )
        assert [action.text() for action in bar.actions() if action.menu()] == expected_menus

        for menu in (window.file_menu, window.help_menu):
            rect = bar.actionGeometry(menu.menuAction())
            for x in (rect.left() + 2, rect.center().x(), rect.right() - 2):
                point = bar.mapToGlobal(QPoint(x, rect.center().y()))
                receiver = QApplication.widgetAt(point)
                assert receiver is bar
                QTest.mouseClick(receiver, Qt.MouseButton.LeftButton, pos=receiver.mapFromGlobal(point))
                qtbot.waitUntil(menu.isVisible)
                menu.hide()
                QApplication.processEvents()

        assert window.new_action.isEnabled()
        assert window.open_action.isEnabled()
        assert not window.project_settings_action.isEnabled()

        about_bar_rect = bar.actionGeometry(window.help_menu.menuAction())
        about_bar_point = bar.mapToGlobal(about_bar_rect.center())
        QTest.mouseClick(bar, Qt.MouseButton.LeftButton, pos=bar.mapFromGlobal(about_bar_point))
        qtbot.waitUntil(window.help_menu.isVisible)
        about_rect = window.help_menu.actionGeometry(window.about_action)
        about_point = window.help_menu.mapToGlobal(about_rect.center())
        about_receiver = QApplication.widgetAt(about_point)
        assert about_receiver is window.help_menu
        about_opened = []

        def close_about() -> None:
            dialog = QApplication.activeModalWidget()
            about_opened.append(dialog is not None and dialog.objectName() == "aboutDialog")
            if dialog is not None:
                dialog.reject()

        QTimer.singleShot(0, close_about)
        QTest.mouseClick(
            about_receiver,
            Qt.MouseButton.LeftButton,
            pos=about_receiver.mapFromGlobal(about_point),
        )
        assert about_opened == [True]
        assert window.help_menu.title() == help_title
        assert not hasattr(window, "about_menu")

        root = tmp_path / "game"
        shutil.copytree(DEMO, root)
        service = project_service()
        window.load_session(service.open_session(root))
        assert window.left_panel.isHidden()
        assert window.icon_toolbar.isVisible()
        assert window.project_settings_action.isEnabled()
        window.ui_state.save(window)

        restored = MainWindow(ProjectController(project_service()))
        qtbot.addWidget(restored)
        assert restored.left_panel.isHidden()
        restored.load_session(project_service().open_session(root))
        assert restored.left_panel.isHidden()
    finally:
        set_language(previous_language)


def test_copy_mode_switch_changes_label_color_and_keyboard_state(qtbot) -> None:
    from componentpress.editor.copy_mode_switch import CopyModeSwitch

    switch = CopyModeSwitch()
    qtbot.addWidget(switch)
    switch.show()
    assert switch.mode == "prod" and switch.text() == "Полный набор (prod)"
    assert switch.accessibleName() == "Режим тиража"
    assert switch.accessibleDescription() == "Режим тиража: Полный набор (prod)"

    def pixel(x: int, y: int) -> tuple[int, int, int]:
        image = QImage(switch.size(), QImage.Format.Format_ARGB32)
        image.fill(Qt.GlobalColor.transparent)
        switch.render(image)
        color = image.pixelColor(x, y)
        return color.red(), color.green(), color.blue()

    assert pixel(25, 15) == (197, 143, 143)
    assert pixel(13, 15) == (255, 255, 255)
    QTest.mouseClick(switch, Qt.MouseButton.LeftButton, pos=switch.rect().center())
    assert switch.mode == "test" and switch.text() == "Частичный набор (test)"
    assert switch.accessibleDescription() == "Режим тиража: Частичный набор (test)"
    assert pixel(10, 15) == (143, 181, 154)
    assert pixel(38, 15) == (255, 255, 255)
    QTest.keyClick(switch, Qt.Key.Key_Space)
    assert switch.mode == "prod" and switch.text() == "Полный набор (prod)"


def test_copy_mode_info_uses_tooltip_without_changing_mode_or_preview(tmp_path: Path, qtbot) -> None:
    _root, window, _service = make_window(tmp_path, qtbot)
    assert window.open_component("forest-card") is not None
    mode = window.data_mode.mode
    generation = window.active_tab.preview_generation
    point = window.mode_info_button.mapToGlobal(QPoint(1, 1))
    QApplication.sendEvent(
        window.mode_info_button,
        QHelpEvent(QEvent.Type.ToolTip, QPoint(1, 1), point),
    )
    assert window._mode_info_tooltip.text() == window.mode_info_button.accessibleName()

    QTest.keyClick(window.mode_info_button, Qt.Key.Key_Space)

    assert "Prod" in window._mode_info_tooltip.text()
    assert "Debug" in window._mode_info_tooltip.text()
    assert window.data_mode.mode == mode
    assert window.active_tab.preview_generation == generation


def test_events_hide_project_path_and_copy_visible_selection(tmp_path: Path, qtbot, monkeypatch) -> None:
    from componentpress.domain.diagnostics import Diagnostic
    from componentpress.i18n.diagnostics import format_diagnostic

    root, window, _service = make_window(tmp_path, qtbot)
    shown = []
    monkeypatch.setattr(main_window_module.QToolTip, "showText", staticmethod(lambda _point, text, *_args: shown.append(text)))
    diagnostic = Diagnostic(
        "RESOURCE_MISSING", "Файл не найден", root / "components" / "forest-card.yaml",
        line=7, node_id="picture",
    )
    window._add_events((diagnostic,))
    item = window.event_list.item(0)
    original_text = item.text()

    assert window.hide_event_project_path.isChecked()
    assert "forest-card.yaml" in original_text
    assert str(root) not in original_text
    assert item.toolTip() == format_diagnostic(diagnostic)
    assert item.data(Qt.ItemDataRole.UserRole) == (diagnostic, root)

    window.hide_event_project_path.setChecked(False)
    assert str(root) in item.text()
    window.hide_event_project_path.setChecked(True)
    window.event_list.setCurrentItem(item)
    window.copy_event.click()
    assert QApplication.clipboard().text() == item.text()
    assert shown == ["Строка события скопирована"]

    window.event_search.setText("строка отсутствует")
    assert item.isHidden() and not window.copy_event.isEnabled()
    window.event_search.clear()
    assert not item.isHidden() and window.copy_event.isEnabled()
    window.clear_events.click()
    assert window.event_list.count() == 0 and not window.copy_event.isEnabled()


@pytest.mark.parametrize("mode", ("prod", "test"))
def test_build_request_uses_copy_switch_mode(tmp_path: Path, qtbot, monkeypatch, mode: str) -> None:
    from types import SimpleNamespace
    from componentpress.application.contracts import BuildRequest
    from componentpress.domain.project import PrintSettings
    from componentpress.user_preferences import ProgramPreferences

    _root, window, _service = make_window(tmp_path, qtbot)
    window.data_mode.setChecked(mode == "test")
    window.save_all = lambda: True
    window.controller.build = SimpleNamespace(cancel=lambda _job_id: True)
    requests = []

    window._program_preferences = ProgramPreferences(
        print_settings=PrintSettings(cut_lines=False, dpi=450), png_workers=5,
    )
    monkeypatch.setattr(window.controller, "start_build", lambda request, **_kwargs: (requests.append(request), "job")[1])
    window._start_build(active_only=False)
    window._build_progress.start_button.click()

    assert isinstance(requests[0], BuildRequest)
    assert requests[0].mode == mode
    assert requests[0].print_settings == PrintSettings(cut_lines=False, dpi=450)
    assert requests[0].max_render_workers == 5
    assert window._build_progress is not None
    assert not window.build_active_widget_action.isVisible()
    assert not window.build_all_widget_action.isVisible()
    assert window.cancel_build_widget_action.isVisible()
    def show_build_error(*_args, **_kwargs):
        assert window.build_active_widget_action.isVisible()
        assert window.build_all_widget_action.isVisible()
        assert not window.cancel_build_widget_action.isVisible()
        return QMessageBox.StandardButton.Ok

    from componentpress.application.contracts import BuildResult
    window._on_build_finished(BuildResult("job", "cancelled"))
    assert window.build_active_widget_action.isVisible()
    assert window.build_all_widget_action.isVisible()
    assert not window.cancel_build_widget_action.isVisible()


@pytest.mark.parametrize("mode", ("prod", "test"))
def test_instance_export_uses_copy_switch_mode(tmp_path: Path, qtbot, monkeypatch, mode: str) -> None:
    from types import SimpleNamespace

    _root, window, _service = make_window(tmp_path, qtbot)
    window.open_component("forest-card")
    window.data_mode.setChecked(mode == "test")
    calls = []
    monkeypatch.setattr(main_window_module.QFileDialog, "getSaveFileName", lambda *_args: (str(tmp_path / "instance.png"), ""))
    window.controller.preview = SimpleNamespace(
        export_png=lambda *_args, **kwargs: (calls.append(kwargs["mode"]), SimpleNamespace(png=tmp_path / "instance.png"))[1]
    )

    window._export_instance()
    assert calls == [mode]


@pytest.mark.parametrize("mode", ("prod", "test"))
@pytest.mark.parametrize("language", ("ru", "en"))
def test_png_zip_uses_copy_switch_mode_and_filename(tmp_path: Path, qtbot, monkeypatch, mode: str, language: str) -> None:
    from types import SimpleNamespace
    from componentpress.i18n import set_language, tr

    set_language(language)
    _root, window, _service = make_window(tmp_path, qtbot)
    window.open_component("forest-card")
    window.data_mode.setChecked(mode == "test")
    window.save_all = lambda: True
    calls = []
    suggested_names = []
    dialog_text = []

    class Choice:
        class ButtonRole:
            AcceptRole = 1

        class StandardButton:
            Cancel = 2

        def __init__(self, _parent):
            self.active = None

        def setWindowTitle(self, value): dialog_text.append(value)
        def setText(self, value): dialog_text.append(value)
        def addButton(self, _value, _role=None):
            if isinstance(_value, str):
                dialog_text.append(_value)
            button = object()
            if self.active is None:
                self.active = button
            return button
        def exec(self): pass
        def clickedButton(self): return self.active

    monkeypatch.setattr(main_window_module, "QMessageBox", Choice)
    def choose_archive(_parent, title, suggested, _filter):
        dialog_text.append(title)
        suggested_names.append(suggested)
        return str(tmp_path / "out.zip"), ""

    monkeypatch.setattr(main_window_module.QFileDialog, "getSaveFileName", choose_archive)
    monkeypatch.setattr(
        window.controller, "png_archive_exporter",
        lambda _snapshot, path, **kwargs: (
            calls.append((Path(path), kwargs["mode"])),
            SimpleNamespace(path=Path(path), png_count=1, copies=1, mode=kwargs["mode"], diagnostics=()),
        )[1],
    )

    window._export_png_archive()
    qtbot.waitUntil(lambda: bool(calls), timeout=5000)
    assert calls == [(tmp_path / "out.zip", mode)]
    assert suggested_names == [f"forest-card-{mode}.zip"]
    assert dialog_text == [
        tr("ui.png_export_title"), tr("ui.png_export_choice"),
        tr("ui.png_active"), tr("ui.png_all"), tr("ui.png_save"),
    ]
    assert window.statusBar().currentMessage().startswith(
        tr("ui.png_created", path=tmp_path / "out.zip", png_count=1, copies=1, mode=mode)
    )


def test_project_settings_action_tracks_project_and_tab_state(tmp_path: Path, qtbot) -> None:
    QSettings.setDefaultFormat(QSettings.Format.IniFormat)
    QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, str(tmp_path / "qsettings"))
    root = tmp_path / "game"
    shutil.copytree(DEMO, root)
    service = project_service()
    service.register(root, "second", "Второй компонент")
    window = MainWindow(ProjectController(service))
    use_temp_ui_settings(window, tmp_path)
    qtbot.addWidget(window)
    assert not window.project_settings_action.isEnabled()
    assert not window.edit_menu.menuAction().isVisible()
    assert not window.view_menu.menuAction().isVisible()
    assert not window.build_menu.menuAction().isVisible()
    assert window.import_project_archive_action.isVisible()
    window.load_session(service.open_session(root))
    assert window.project_settings_action.isEnabled()  # A project can have no open document.
    window.session.documents.clear()
    window._update_actions()
    assert window.project_settings_action.isEnabled()  # Empty projects still have project settings.
    window.controller.set_session(None)
    window._update_actions()
    assert not window.project_settings_action.isEnabled()
    assert not window.project_settings_action.isVisible()
    assert not window.edit_menu.menuAction().isVisible()
    assert not window.view_menu.menuAction().isVisible()
    assert not window.build_menu.menuAction().isVisible()
    assert window.import_project_archive_action.isVisible()


def test_window_menu_actions_toggle_panels_and_sync_close_button(tmp_path: Path, qtbot) -> None:
    root, window, service = make_window(tmp_path, qtbot)
    actions = window.windows_menu.actions()
    assert [action.text() for action in actions] == ["Рабочие панели", "События"]
    assert all(action.isCheckable() for action in actions)
    window.controller.set_session(None)
    window._update_project_visibility()
    assert all(not action.isEnabled() for action in actions)
    window.load_session(service.open_session(root))
    assert all(action.isEnabled() for action in actions)

    for action, dock, area in zip(actions, window.docks, (
        Qt.DockWidgetArea.RightDockWidgetArea, Qt.DockWidgetArea.BottomDockWidgetArea,
    )):
        dock.hide()
        action.trigger()
        assert dock.isVisible() and action.isChecked()
        assert window.dockWidgetArea(dock) == area
        action.trigger()
        assert not dock.isVisible() and not action.isChecked()
        dock.show()
        qtbot.waitUntil(action.isChecked)
        dock.hide()
        assert not action.isChecked()

    window.events_dock.hide()
    window.ui_state.save(window)
    window._update_project_visibility()
    restored = MainWindow(ProjectController(project_service()))
    restored.ui_state = UiState(QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat))
    restored.ui_state.restore(restored)
    qtbot.addWidget(restored)
    restored.load_session(project_service().open_session(root))
    restored.show()
    assert not restored.events_dock.isVisible()
    assert not hasattr(restored, "toolbox_dock")
    assert not hasattr(restored, "preview_dock")


def test_quick_settings_action_opens_project_settings(tmp_path: Path, qtbot, monkeypatch) -> None:
    _root, window, _service = make_window(tmp_path, qtbot)
    opened = []
    monkeypatch.setattr(ProjectSettingsDialog, "exec", lambda dialog: opened.append(dialog._project_root))

    window.project_settings_action.trigger()

    assert opened == [window.session.snapshot.root]


def test_menu_popup_renders_native_window_contents_with_effects_disabled(tmp_path: Path, qtbot) -> None:
    _root, window, _service = make_window(tmp_path, qtbot)
    screen = window.screen()

    assert not screen is None
    assert not window.windowHandle() is None
    assert not QApplication.isEffectEnabled(Qt.UIEffect.UI_AnimateMenu)
    assert not QApplication.isEffectEnabled(Qt.UIEffect.UI_FadeMenu)
    menu = window.file_menu
    menu.popup(window.menuBar().mapToGlobal(window.menuBar().actionGeometry(menu.menuAction()).bottomLeft()))
    qtbot.waitUntil(menu.isVisible, timeout=1000)
    capture = menu.grab().toImage()
    assert not capture.isNull()
    colors = {capture.pixelColor(x, y).rgba() for y in range(capture.height()) for x in range(capture.width())}
    assert len(colors) >= 4
    menu.hide()

def test_delete_component_action_requires_active_tab_and_deletes_selected_component(tmp_path: Path, qtbot, monkeypatch) -> None:
    root, window, service = make_window(tmp_path, qtbot)
    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *args, **kwargs: QMessageBox.StandardButton.Yes))
    assert not window.delete_component_action.isEnabled()
    window.load_session(service.open_session(root))
    assert not window.delete_component_action.isEnabled()  # Project open, no active tab.
    first = window.open_component("forest-card")
    assert first and window.delete_component_action.isEnabled()
    window._request_mode(first, "text")
    assert not first.text_editor.document().isUndoAvailable()
    second = window.open_component("second")
    window.tabs.setCurrentWidget(second)
    assert window.active_tab is second and window.active_tab.component_id == "second"
    window.delete_component_action.trigger()
    assert "second" not in window.session.documents
    assert not window.delete_component_action.isEnabled()  # Only one component remains.
    window._undo()
    assert "second" in window.session.documents
    window._redo()
    assert "second" not in window.session.documents
    window._undo()
    assert "second" in window.session.documents
    for index in reversed(range(window.tabs.count())):
        window.close_tab(index)
    assert not window.delete_component_action.isEnabled()


def test_closing_restored_component_before_branching_history_is_safe(tmp_path: Path, qtbot) -> None:
    _root, window, _service = make_window(tmp_path, qtbot)
    tab = window.open_component("second")
    assert tab is not None
    assert window.delete_component("second", confirm=False)
    window._undo()
    assert window._tabs["second"] is tab
    assert window.close_tab(window.tabs.indexOf(tab))
    qtbot.wait(20)
    assert window.add_component("third", "Третий компонент")
    assert "third" in window.session.documents


def test_excel_picker_preserves_nested_project_path_and_reuses_workbook(tmp_path: Path, qtbot, monkeypatch) -> None:
    source_root = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"
    root = tmp_path / "game"
    shutil.copytree(source_root, root)
    service = project_service()
    events = []
    service.event_logger = lambda event, **context: events.append((event, context))
    window = MainWindow(ProjectController(service))
    use_temp_ui_settings(window, tmp_path)
    window.load_session(service.open_session(root))
    qtbot.addWidget(window)
    events.clear()
    workbook_before = (root / "data/game.xlsx").read_bytes()
    dialog = ProjectSettingsDialog(window.session, window.controller.preview, window)
    dialog.version.setText("34.0")
    monkeypatch.setattr(QFileDialog, "getOpenFileName", staticmethod(lambda *a, **kw: (str(root / "data/game.xlsx"), "")))
    dialog._choose_file()
    assert dialog.source.text() == "data/game.xlsx"
    window.controller.update_project_settings(
        name=dialog.name.text(), version=dialog.version.text(), prod=dialog.prod.text(), test=dialog.test.text(),
        data_source_file=dialog.values()["source_file"],
    )
    assert window.session.snapshot.model.data_source.path == "data/game.xlsx"
    assert window.session.snapshot.model.version == "34.0"
    assert events[-1][1]["status"] == "succeeded"
    reopened = FileProjectRepository().open(root)
    assert reopened.model.data_source.path == "data/game.xlsx"
    assert (root / "data/game.xlsx").read_bytes() == workbook_before
    assert sorted(path.name for path in (root / "data").glob("*.xlsx")) == ["game.xlsx"]


@pytest.mark.parametrize("locale", ("ru", "en"))
def test_invalid_version_is_localized_and_rejected_before_excel_copy(tmp_path: Path, qtbot, locale: str) -> None:
    from componentpress.i18n import set_language, tr

    set_language(locale)
    source_root = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"
    root = tmp_path / "game"
    shutil.copytree(source_root, root)
    external = tmp_path / "outside.xlsx"
    shutil.copy2(root / "data/game.xlsx", external)
    service = project_service()
    events = []
    service.event_logger = lambda event, **context: events.append((event, context))
    window = MainWindow(ProjectController(service))
    use_temp_ui_settings(window, tmp_path)
    window.load_session(service.open_session(root))
    qtbot.addWidget(window)
    events.clear()
    before = (root / "project.cpress").read_bytes()

    from componentpress.domain.diagnostics import ProjectError

    with pytest.raises(ProjectError) as caught:
        service.update_project_settings(
            window.session, name="Game", version="toolong", prod="Prod", test="Debug",
            data_source_file=external,
        )
    assert caught.value.diagnostic.code == "PROJECT_VERSION_INVALID"
    assert tr("error.project_version_invalid") in format_error(caught.value)
    assert (root / "project.cpress").read_bytes() == before
    assert not (root / "data/outside.xlsx").exists()
    assert events[0][1] == {"operation": "project_settings", "status": "started"}
    assert events[-1][1]["code"] == "PROJECT_VERSION_INVALID"


def test_project_settings_external_nested_book_and_conflict_are_safe(tmp_path: Path, qtbot) -> None:
    source_root = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"
    root = tmp_path / "game"
    shutil.copytree(source_root, root)
    external = tmp_path / "outside" / "deep" / "book.xlsx"
    external.parent.mkdir(parents=True)
    shutil.copy2(root / "data/game.xlsx", external)
    service = project_service()
    window = MainWindow(ProjectController(service))
    use_temp_ui_settings(window, tmp_path)
    window.load_session(service.open_session(root))
    qtbot.addWidget(window)
    window.controller.update_project_settings(name="Game", version="0.1.0", prod="Prod", test="Debug", data_source_file=external)
    assert window.session.snapshot.model.data_source.path == "data/book.xlsx"
    assert (root / "data/book.xlsx").read_bytes() == external.read_bytes()
    conflict = tmp_path / "conflict" / "book.xlsx"
    conflict.parent.mkdir()
    conflict.write_bytes(b"not an xlsx")
    errors = []
    window._show_error = errors.append
    before_project = (root / "project.cpress").read_bytes()
    before_book = (root / "data/book.xlsx").read_bytes()
    try:
        window.controller.update_project_settings(name="Game", version="0.1.0", prod="Prod", test="Debug", data_source_file=conflict)
    except Exception as exc:
        assert str(exc).strip()
    else:
        raise AssertionError("invalid external XLSX unexpectedly accepted")
    assert (root / "project.cpress").read_bytes() == before_project
    assert (root / "data/book.xlsx").read_bytes() == before_book


def test_import_cleanup_hash_failure_preserves_primary_error_and_reused_book(tmp_path: Path, monkeypatch) -> None:
    source_root = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"
    root = tmp_path / "game"
    shutil.copytree(source_root, root)
    service = project_service()
    session = service.open_session(root)
    source = tmp_path / "new.xlsx"
    shutil.copy2(root / "data/game.xlsx", source)
    original_book = (root / "data/game.xlsx").read_bytes()
    project_before = (root / "project.cpress").read_bytes()
    primary_error = RuntimeError("save failure must stay primary")
    monkeypatch.setattr(service.projects, "save_project", lambda *_args, **_kwargs: (_ for _ in ()).throw(primary_error))

    from componentpress.resources import repository as resource_repository

    real_hash = resource_repository.file_hash
    calls = 0

    def fail_rollback_hash(path: Path) -> str:
        nonlocal calls
        calls += 1
        if calls == 1:
            return real_hash(path)
        raise OSError("cleanup hash failed")

    monkeypatch.setattr(resource_repository, "file_hash", fail_rollback_hash)
    try:
        service.update_project_settings(
            session, name=session.snapshot.model.name, version=session.snapshot.model.version,
            prod="Prod", test="Debug", data_source_file=source,
        )
    except RuntimeError as exc:
        assert exc is primary_error
    else:
        raise AssertionError("primary save error was not raised")

    assert calls == 2
    assert (root / "data/game.xlsx").read_bytes() == original_book
    assert (root / "project.cpress").read_bytes() == project_before


def test_blank_errors_and_invalid_project_version_have_useful_message(tmp_path: Path, qtbot, monkeypatch) -> None:
    assert format_error(ValueError(" \n "))
    root, window, service = make_window(tmp_path, qtbot)
    window.load_session(service.open_session(root))
    errors = []
    original_show_error = window._show_error
    monkeypatch.setattr(window, "_show_error", errors.append)
    class InvalidSettingsDialog:
        DialogCode = QDialog.DialogCode

        def __init__(self, *_args):
            self.calls = 0

        def exec(self):
            self.calls += 1
            return QDialog.DialogCode.Accepted if self.calls == 1 else QDialog.DialogCode.Rejected

        def values(self):
            return {"name": "Game", "version": "too long", "prod": "Prod", "test": "Debug", "source": "", "source_file": None}

    instances = []

    def make_invalid_dialog(*args):
        dialog = InvalidSettingsDialog(*args)
        instances.append(dialog)
        return dialog

    monkeypatch.setattr(main_window_module, "ProjectSettingsDialog", make_invalid_dialog)
    window.project_settings_dialog()
    assert instances[0].calls == 2
    assert window.session.snapshot.model.version != "too long"
    assert errors and errors[0].diagnostic.code == "PROJECT_VERSION_INVALID"

    shown = []
    monkeypatch.setattr(QMessageBox, "critical", staticmethod(lambda _parent, _title, message: shown.append(message)))
    monkeypatch.setattr(window, "_show_error", original_show_error)
    window._show_error(ValueError("  \n"))
    assert shown and shown[0].strip()


def test_demo_excel_component_tree_is_the_visible_interactive_tree(tmp_path: Path, qtbot) -> None:
    source_root = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"
    root = tmp_path / "demo-excel-game"
    shutil.copytree(source_root, root)
    service = project_service()
    window = MainWindow(ProjectController(service))
    use_temp_ui_settings(window, tmp_path)
    window.load_session(service.open_session(root))
    qtbot.addWidget(window)
    window.show()
    tree = window.right_tabs.widget(0).findChild(QTreeView, "componentTree")
    assert tree is window.component_tree and tree.model().rowCount() == 2
    ids = [tree.model().index(row, 0).data(ID_ROLE) for row in range(tree.model().rowCount())]
    assert ids == ["forest-card", "event-card"]
    event_index = tree.model().index(1, 0)
    qtbot.mouseClick(tree.viewport(), Qt.MouseButton.LeftButton, pos=tree.visualRect(event_index).center())
    qtbot.waitUntil(lambda: window.active_tab is not None and window.active_tab.component_id == "event-card")
    forest = tree.model().index(0, 0)
    qtbot.mouseClick(tree.viewport(), Qt.MouseButton.LeftButton, pos=tree.visualRect(forest).center())
    qtbot.waitUntil(lambda: window.active_tab.component_id == "forest-card")
    tree.setFocus()
    QTest.keyClick(tree, Qt.Key.Key_Down)
    qtbot.waitUntil(lambda: window.active_tab.component_id == "event-card")
    other = window.open_component("event-card")
    assert other is not None
    assert tree.currentIndex().data(ID_ROLE) == "event-card"
    forest_tab = window.open_component("forest-card")
    window.tabs.setCurrentWidget(forest_tab)
    qtbot.waitUntil(lambda: tree.currentIndex().data(ID_ROLE) == "forest-card")
    while window.tabs.count():
        window.close_tab(window.tabs.count() - 1)
    assert not tree.currentIndex().isValid()


def test_loading_excel_project_validates_every_component_and_renders_selected_layout(tmp_path: Path, qtbot) -> None:
    from componentpress.bootstrap import project_controller

    source_root = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"
    root = tmp_path / "demo-excel-game"
    shutil.copytree(source_root, root)
    window = MainWindow(project_controller())
    use_temp_ui_settings(window, tmp_path)
    qtbot.addWidget(window)
    window.load_session(window.controller.open(root))

    snapshot = window.session.snapshot
    assert window.controller.preview.cached_data(snapshot, "forest-card") is not None
    assert window.controller.preview.cached_data(snapshot, "event-card") is not None
    tab = window.open_component("forest-card")
    assert tab is not None
    qtbot.waitUntil(lambda: tab.displayed_generation == tab.preview_generation)
    assert tab.canvas._pixmap_item is not None
    assert {"leaf-art", "text-group", "heading", "caption"}.issubset(tab.canvas._overlays)


def test_moving_and_undoing_element_keeps_excel_preview_visible(tmp_path: Path, qtbot) -> None:
    from componentpress.bootstrap import project_controller

    root = tmp_path / "demo-excel-game"
    shutil.copytree(Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game", root)
    window = MainWindow(project_controller())
    use_temp_ui_settings(window, tmp_path)
    qtbot.addWidget(window)
    window.load_session(window.controller.open(root))
    tab = window.open_component("forest-card")
    assert tab is not None
    qtbot.waitUntil(lambda: tab.displayed_generation == tab.preview_generation)
    service = window.controller.preview
    snapshot = window.session.snapshot
    data = service.cached_data(snapshot, "forest-card")
    assert data is not None and tab.canvas._pixmap_item is not None

    previous_item = tab.canvas._pixmap_item
    window._move_selected(("leaf-art",), 2.0, 0.0)
    assert tab.canvas._pixmap_item.isVisible()
    qtbot.waitUntil(lambda: tab.displayed_generation == tab.preview_generation)
    assert tab.canvas._pixmap_item.isVisible()
    assert service.cached_data(snapshot, "forest-card") is data
    assert tab.canvas._pixmap_item is not None
    assert tab.canvas._pixmap_item is previous_item

    tab.undo_stack.undo()
    assert tab.canvas._pixmap_item.isVisible()
    qtbot.waitUntil(lambda: tab.displayed_generation == tab.preview_generation)
    assert tab.canvas._pixmap_item.isVisible()
    assert service.cached_data(snapshot, "forest-card") is data
    assert tab.canvas._pixmap_item is not None


def test_loading_excel_project_reports_bad_component_and_keeps_other_validation(tmp_path: Path, qtbot) -> None:
    from openpyxl import load_workbook
    from componentpress.bootstrap import project_controller

    source_root = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"
    root = tmp_path / "demo-excel-game"
    shutil.copytree(source_root, root)
    book_path = root / "data" / "game.xlsx"
    book = load_workbook(book_path)
    book["Карты"]["D2"] = "missing.png"
    book.save(book_path)
    book.close()

    window = MainWindow(project_controller())
    use_temp_ui_settings(window, tmp_path)
    qtbot.addWidget(window)
    window.load_session(window.controller.open(root))
    snapshot = window.session.snapshot

    assert window.controller.preview.cached_data(snapshot, "forest-card") is None
    assert window.controller.preview.cached_data(snapshot, "event-card") is not None
    assert any("RESOURCE_MISSING" in window.event_list.item(i).text() for i in range(window.event_list.count()))
    assert not window.events_dock.isHidden()
    tab = window.open_component("forest-card")
    assert tab is not None
    assert not tab.preview_first.isEnabled()
    assert not tab.preview_next.isEnabled()


def test_loading_excel_project_with_no_prod_instances_disables_instance_navigation(tmp_path: Path, qtbot) -> None:
    from openpyxl import load_workbook
    from componentpress.bootstrap import project_controller

    source_root = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"
    root = tmp_path / "demo-excel-game"
    shutil.copytree(source_root, root)
    book_path = root / "data" / "game.xlsx"
    book = load_workbook(book_path)
    book["Карты"]["I2"] = 0
    book["Карты"]["I3"] = 0
    book.save(book_path)
    book.close()

    window = MainWindow(project_controller())
    use_temp_ui_settings(window, tmp_path)
    qtbot.addWidget(window)
    window.load_session(window.controller.open(root))
    tab = window.open_component("forest-card")

    assert tab is not None
    assert not tab.preview_first.isEnabled()
    assert not tab.preview_next.isEnabled()
    assert window.controller.preview.cached_data(window.session.snapshot, "forest-card") is not None
    assert not any("VALIDATION_ERROR" in window.event_list.item(i).text() for i in range(window.event_list.count()))


def test_reopening_project_invalidates_previous_excel_validation(tmp_path: Path, qtbot) -> None:
    from openpyxl import load_workbook
    from componentpress.bootstrap import project_controller

    source_root = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"
    root = tmp_path / "demo-excel-game"
    shutil.copytree(source_root, root)
    controller = project_controller()
    window = MainWindow(controller)
    use_temp_ui_settings(window, tmp_path)
    qtbot.addWidget(window)
    window.load_session(controller.open(root))
    assert controller.preview.cached_data(window.session.snapshot, "forest-card") is not None

    book_path = root / "data" / "game.xlsx"
    book = load_workbook(book_path)
    book["Карты"]["D2"] = "missing.png"
    book.save(book_path)
    book.close()
    window.load_session(controller.open(root))

    assert controller.preview.cached_data(window.session.snapshot, "forest-card") is None
    assert controller.preview.cached_data(window.session.snapshot, "event-card") is not None
    assert any("RESOURCE_MISSING" in window.event_list.item(i).text() for i in range(window.event_list.count()))


def test_compact_shell_has_icon_tools_and_project_docks(tmp_path: Path, qtbot) -> None:
    _root, window, _service = make_window(tmp_path, qtbot)

    assert window.right_tabs.tabText(0) == "Компоненты"
    assert window.right_tabs.tabText(1) == "Элементы"
    assert window.events_dock.isVisible()
    assert len(window.icon_toolbar.actions()) == 14
    assert sum(action.isSeparator() for action in window.icon_toolbar.actions()) == 2
    assert window.project_settings_action in window.quick_toolbar.actions()
    assert window.project_settings_action not in window.icon_toolbar.actions()
    assert all(action.text() not in {"Инструменты", "Превью"} for action in window.windows_menu.actions())
    assert not hasattr(window, "toolbox_dock")
    assert not hasattr(window, "preview_dock")
    assert window.select_tool.icon().isNull() is False
    assert window.select_tool.text() == ""
    assert window.select_tool.toolTip() == "Выбор"
    assert window.help_menu.title() == "Справка"


def test_command_and_instance_buttons_use_the_assigned_png_icons(tmp_path: Path, qtbot, monkeypatch) -> None:
    from componentpress.user_preferences import ProgramPreferences
    monkeypatch.setattr(main_window_module, "load_preferences", lambda: ProgramPreferences(theme="light"))
    _root, window, _service = make_window(tmp_path, qtbot)
    assets = Path(__file__).resolve().parents[1] / "src" / "componentpress" / "app_assets"

    def matches_asset(icon: QIcon, filename: str) -> bool:
        size = QSize(22, 22)
        expected = QIcon(QPixmap(str(assets / filename))).pixmap(size, QIcon.Mode.Normal, QIcon.State.Off)
        actual = icon.pixmap(size, QIcon.Mode.Normal, QIcon.State.Off)
        return not actual.isNull() and actual.toImage() == expected.toImage()

    action_icons = {
        window.project_settings_action: "settings.png", window.add_component_action: "add.png",
        window.save_action: "save.png", window.undo_action: "undo.png", window.redo_action: "redo.png",
        window.delete_action: "remove.png",
    }
    for action, filename in action_icons.items():
        assert matches_asset(action.icon(), filename), (action.text(), filename)
        assert action.toolTip() == action.text()
        assert action.property("accessibleName") == action.text()

    assert window.actual_size_action not in window.icon_toolbar.actions()
    assert window.validate_active_action not in window.quick_toolbar.actions()
    assert window.validate_all_action.text() == "Обновить данные и проверить проект"
    assert matches_asset(window.validate_all_action.icon(), "validate.png")
    assert matches_asset(window.build_active_action.icon(), "build.png")
    assert matches_asset(window.build_all_action.icon(), "build_all.png")
    assert matches_asset(window.cancel_build_action.icon(), "stop.png")
    assert not window.cancel_build_widget_action.isVisible()
    assert not window.delete_widget_action.isVisible()
    window.open_component("forest-card")
    window._set_selection(("heading",), source="test")
    assert window.delete_widget_action.isVisible()
    window._set_selection((), source="test")
    assert not window.delete_widget_action.isVisible()
    window.active_tab.canvas.zoom_by(1.25)
    assert window.actual_size_action in window.actions()
    assert window.actual_size_action.shortcut().toString() == "Ctrl+1"
    window.actual_size_action.trigger()
    assert window.active_tab.canvas.zoom_percent == 100
    window._active_build_job = "job"
    window._update_actions()
    assert not window.build_active_widget_action.isVisible()
    assert not window.build_all_widget_action.isVisible()
    assert window.cancel_build_widget_action.isVisible()
    window._active_build_job = None
    window._version_task = object()
    window._update_actions()
    assert window.build_active_widget_action.isVisible()
    assert window.build_all_widget_action.isVisible()
    assert not window.cancel_build_widget_action.isVisible()
    assert not window.cancel_build_action.isEnabled()
    window._version_task = None
    window._update_actions()

    widget_icons = {
        window.select_tool: "cursor.png", window.image_tool: "image.png", window.html_tool: "text.png",
        window.conditional_group_tool: "condition.png", window.rectangle_tool: "rectangle.png",
        window.ellipse_tool: "ellipse.png", window.line_tool: "line.png", window.layer_up: "up.png",
        window.layer_down: "down.png", window.group_button: "group.png", window.ungroup_button: "ungroup.png",
        window.clear_events: "delete.png",
        window.add_component_button: "add.png",
    }
    for button, filename in widget_icons.items():
        assert matches_asset(button.icon(), filename), (button.accessibleName(), filename)
    assert window.clear_events.text() == "Очистить"
    assert window.group_button.accessibleName() == "Сгруппировать"
    assert not hasattr(window, "reparent_button")
    assert window.add_component_button.text() == ""

    tab = window.open_component("forest-card")
    assert tab is not None
    instance_icons = (
        (tab.preview_first, "begin.png", "Первый экземпляр"),
        (tab.preview_previous, "left.png", "Предыдущий"),
        (tab.preview_next, "right.png", "Следующий"),
        (tab.preview_last, "end.png", "Последний экземпляр"),
    )
    for button, filename, label in instance_icons:
        assert matches_asset(button.icon(), filename)
        assert button.text() == ""
        assert button.toolTip() == label
        assert button.accessibleName() == label
        assert button.iconSize() == QSize(22, 22)
    assert not tab.preview_first.isHidden()
    tab.set_mode("text")
    assert all(button.isHidden() for button in tab.instance_controls)
    tab.set_mode("layout")
    assert all(not button.isHidden() for button in tab.instance_controls)
    assert tab.instance_position.text() == "Экземпляр 1 из 1"
    assert matches_asset(tab.layout_mode.icon(), "brick.png")
    assert matches_asset(tab.text_mode.icon(), "code.png")

    settings_action = window.project_settings_action
    assert settings_action in window.quick_toolbar.actions()
    assert settings_action.icon().isNull() is False


def test_delete_selected_requires_confirmation_and_is_one_undo(tmp_path: Path, qtbot) -> None:
    from componentpress.domain.tree import locations

    _root, window, _service = make_window(tmp_path, qtbot)
    window.open_component("forest-card")
    document = window.session.documents["forest-card"]
    tab = window.active_tab
    window._set_selection(("heading",), source="test")
    before = document.model
    history_before = tab.undo_stack.count()

    def cancel_delete() -> None:
        dialog = QApplication.activeModalWidget()
        assert dialog is not None
        assert dialog.windowTitle() == "Подтвердите удаление"
        assert "heading" not in dialog.text()
        assert any(button.text() == "Удалить" for button in dialog.buttons())
        next(button for button in dialog.buttons() if button.text() == "Отмена").click()

    QTimer.singleShot(0, cancel_delete)
    assert not window.delete_selected()
    assert document.model == before
    assert tab.undo_stack.count() == history_before
    assert "heading" in locations(document.model)

    QTimer.singleShot(0, lambda: next(
        button for button in QApplication.activeModalWidget().buttons() if button.text() == "Удалить"
    ).click())
    assert window.delete_selected()
    assert "heading" not in locations(document.model)
    assert tab.undo_stack.count() == history_before + 1
    window.undo_group.undo()
    assert "heading" in locations(document.model)


def test_group_confirmation_rejects_a_changed_selection(tmp_path: Path, qtbot, monkeypatch) -> None:
    _root, window, _service = make_window(tmp_path, qtbot)
    window.open_component("forest-card")
    window._set_selection(("heading", "caption"), source="test")
    before = window.session.documents["forest-card"].model

    def change_selection(*_args, **_kwargs):
        window._set_selection(("heading",), source="test")
        return "Group", True

    monkeypatch.setattr(main_window_module.QInputDialog, "getText", change_selection)
    assert not window.group_selected()
    assert window.session.documents["forest-card"].model == before
    assert not any(node.type == "group" and node.name == "Group" for node in before.elements)


@pytest.mark.parametrize("locale", ("ru", "en"))
def test_component_creation_and_deletion_round_trip_preserves_draft_and_tab(tmp_path: Path, qtbot, monkeypatch, locale: str) -> None:
    from componentpress.i18n import set_language

    set_language(locale)
    root, window, service = make_window(tmp_path, qtbot)
    errors = []
    monkeypatch.setattr(window, "_show_error", errors.append)
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *args, **kwargs: QMessageBox.StandardButton.Yes))
    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *args, **kwargs: QMessageBox.StandardButton.Discard))
    monkeypatch.setattr(window, "_ask_unsaved_choice", lambda *_args: QMessageBox.StandardButton.Discard)

    unlink = Path.unlink

    def fail_add_marker_cleanup(path, *args, **kwargs):
        if path.name == "add-component.json":
            raise PermissionError("injected post-commit cleanup failure")
        return unlink(path, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", fail_add_marker_cleanup)
    assert window.add_component("third", "Третий компонент")
    assert len(window._component_history.entries) == 1
    assert window._component_history.index == 1
    monkeypatch.setattr(Path, "unlink", unlink)
    tab = window.active_tab
    document = window.session.documents["third"]
    window._draft_changed(tab, "schema_version: 8\nid: third\nname: [неверный черновик]\n")
    window._undo()
    assert "third" not in window.session.documents
    assert not (root / "components" / "third.yaml").exists()
    window._redo()
    assert window.session.documents["third"] is document
    assert window.active_tab is tab
    assert document.has_invalid_draft
    assert document.draft_text.startswith("schema_version: 8")
    assert service.projects.open(root).documents["third"].model.name == "Третий компонент"
    assert not errors, errors

    assert window.delete_component("third", confirm=False)
    window._undo()
    assert "third" in window.session.documents
    assert window.active_tab is tab
    assert document.has_invalid_draft
    assert service.projects.open(root).documents["third"].model.name == "Третий компонент"
    assert not errors, errors


def test_failed_component_delete_does_not_hide_tab_or_advance_history(tmp_path: Path, qtbot, monkeypatch) -> None:
    root, window, _service = make_window(tmp_path, qtbot)
    tab = window.open_component("forest-card")
    assert window.add_component("branch", "Branch")
    window._undo()
    redo_entry = window._component_history.redo_entry()
    before = window._component_history_index
    (root / "project.cpress").write_text("external change\n", encoding="utf-8")
    errors = []
    monkeypatch.setattr(window, "_show_error", errors.append)

    assert not window.delete_component("forest-card", confirm=False)
    assert window._component_history_index == before
    assert window._component_history.redo_entry() is redo_entry
    assert window.tabs.indexOf(tab) >= 0
    assert (root / "components" / "forest-card.yaml").exists()
    assert errors


def test_ui_layout_preferences_are_local_and_restored(tmp_path: Path, qtbot) -> None:
    root, window, _service = make_window(tmp_path, qtbot)
    settings = QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat)
    state = UiState(settings)
    window.right_tabs.setCurrentIndex(1)
    window.right_splitter.setSizes([160, 440])
    saved_sizes = window.right_splitter.sizes()
    state.save(window)
    settings.sync()

    other = MainWindow(ProjectController(project_service()))
    other.ui_state = state
    other.ui_state.restore(other)
    assert other.right_tabs.currentIndex() == 1
    other.load_session(project_service().open_session(root))
    other.show()
    qtbot.waitUntil(lambda: sum(other.right_splitter.sizes()) > 0)
    restored_sizes = other.right_splitter.sizes()
    assert restored_sizes[0] / sum(restored_sizes) == pytest.approx(
        saved_sizes[0] / sum(saved_sizes), abs=0.03
    )
    assert (root / "project.cpress").exists()
    assert ".componentpress" not in (root / "project.cpress").read_text(encoding="utf-8")


def test_component_dialog_keeps_fractional_existing_size_and_lists_presets(qtbot) -> None:
    component = project_service().projects.open(DEMO).documents["forest-card"].model.model_copy(
        update={"size_mm": project_service().projects.open(DEMO).documents["forest-card"].model.size_mm.model_copy(
            update={"width": 63.25, "height": 88.5}
        )}
    )
    dialog = ComponentDialog(component=component)
    qtbot.addWidget(dialog)

    assert dialog.width.value() == 63.25
    assert dialog.height.value() == 88.5
    assert dialog.definition(component)["width"] == 63.25
    assert dialog.size.count() >= 15
