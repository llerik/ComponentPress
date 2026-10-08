"""Independent acceptance checks for stage 6 Excel preview behavior."""

from __future__ import annotations

from datetime import date
from pathlib import Path
import shutil
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
import pytest
from PySide6.QtGui import QImage

from componentpress.application.contracts import DocumentSnapshot, ProjectSnapshot
from componentpress.application.preview_service import PreviewService
from componentpress.bindings.resolver import BindingResolver
from componentpress.bindings.static import resolve_static
from componentpress.bootstrap import project_repository
from componentpress.cli import main as cli_main
from componentpress.data_sources import XlsxReader
from componentpress.domain.component import ComponentDefinition, DataBinding, SizeMM
from componentpress.domain.diagnostics import ProjectError
from componentpress.domain.nodes import HtmlNode, ImageNode
from componentpress.domain.project import DataSource, IconDefinition, ProjectDefinition
from componentpress.domain.values import CellValue, DataRow, DataSheetSnapshot, FormulaState, value_to_text
from componentpress.editor.main_window import MainWindow
from componentpress.bootstrap import project_controller, project_service


ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "examples" / "demo-excel-game"
STATIC_DEMO = ROOT / "examples" / "demo-game"


def _write_table(path: Path, rows: list[list[object]], *, title: str = "Лист") -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    book.save(path)
    book.close()


def _patch_formula_cells(path: Path) -> None:
    """Create four genuine formula-cache states without asking openpyxl to calculate."""
    replacement = path.with_suffix(".rewritten.xlsx")
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                from xml.etree import ElementTree as ET

                ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                ET.register_namespace("", ns)
                root = ET.fromstring(data)
                cells = {cell.attrib["r"]: cell for cell in root.iter(f"{{{ns}}}c")}
                for coordinate, state in {
                    "B2": "cached",
                    "C2": "empty",
                    "D2": "missing",
                    "E2": "error",
                }.items():
                    cell = cells[coordinate]
                    cached = cell.find(f"{{{ns}}}v")
                    if cached is not None:
                        cell.remove(cached)
                    if state != "missing":
                        cached = ET.SubElement(cell, f"{{{ns}}}v")
                        cached.text = "7" if state == "cached" else "#VALUE!" if state == "error" else None
                    if state == "error":
                        cell.set("t", "e")
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, data)
    replacement.replace(path)


def _project(*, variables: dict[str, str] | None = None) -> ProjectDefinition:
    return ProjectDefinition(
        schema_version=4,
        id="test",
        name="Тест",
        version="0.0.1",
        variables=variables or {},
        data_source=DataSource(path="data/test.xlsx"),
    )


def _html_component(html: str) -> ComponentDefinition:
    return ComponentDefinition(
        schema_version=4,
        id="card",
        name="Карта",
        size_mm=SizeMM(width=40.0, height=30.0),
        data=DataBinding(source="main", sheet="Лист"),
        elements=(HtmlNode(
            id="text", name="Текст", type="html", x_mm=0.0, y_mm=0.0,
            width_mm=40.0, height_mm=30.0, font_family="Arial", font_size_pt=8.0,
            html=html,
        ),),
    )


def _data_row(**values: object) -> tuple[DataSheetSnapshot, DataRow]:
    cells = {
        name: CellValue(value=value, coordinate=f"{chr(65 + index)}2")
        for index, (name, value) in enumerate(values.items())
    }
    row = DataRow(row_number=2, instance_id="2", copies=1, values=cells)
    data = DataSheetSnapshot(
        source="main", path=Path("data/test.xlsx"), sheet="Лист",
        content_hash="hash", version=1, headers=tuple(values), rows=(row,),
    )
    return data, row


def test_reader_trims_headers_keeps_hidden_rows_skips_blank_rows_and_converts_types(tmp_path: Path) -> None:
    path = tmp_path / "таблица.xlsx"
    _write_table(path, [
        [" ID ", "Текст", "Количество", "Дата", "Флаг", "Дробь"],
        ["hidden", "скрытая", 0, date(2026, 10, 3), True, 2.5],
        [None, None, None, None, None, None],
        ["shown", "видимая", 3, None, False, 4.0],
    ])
    book = load_workbook(path)
    book.active.row_dimensions[2].hidden = True
    book.save(path)
    book.close()

    snapshot = XlsxReader().read(path, "Лист", id_column="ID", copies_column="Количество")

    assert snapshot.headers == ("ID", "Текст", "Количество", "Дата", "Флаг", "Дробь")
    assert [(row.row_number, row.instance_id, row.copies) for row in snapshot.rows] == [
        (2, "hidden", 0), (4, "shown", 3),
    ]
    assert value_to_text(snapshot.rows[0].values["Дата"].value) == "2026-10-03"
    assert value_to_text(snapshot.rows[0].values["Флаг"].value) == "true"
    assert value_to_text(snapshot.rows[0].values["Дробь"].value) == "2.5"
    assert value_to_text(snapshot.rows[1].values["Дробь"].value) == "4"

    # The source must be writable/reopenable immediately after the adapter returns.
    reopened = load_workbook(path)
    reopened.active["B4"] = "изменено после чтения"
    reopened.save(path)
    reopened.close()


@pytest.mark.parametrize(
    ("headers", "rows", "kwargs", "code", "cell"),
    [
        (["ID", " ID "], [["a", "b"]], {}, "HEADER_DUPLICATE", "B1"),
        (["ID", "На{звание}"], [["a", "b"]], {}, "HEADER_BRACES", "B1"),
        (["ID", 42], [["a", "b"]], {}, "HEADER_INVALID", "B1"),
        (["ID", "Количество"], [["", 1]], {"id_column": "ID"}, "INSTANCE_ID_EMPTY", "A2"),
        (["ID", "Количество"], [["a", 1], ["a", 2]], {"id_column": "ID"}, "INSTANCE_ID_DUPLICATE", "A3"),
        (["ID", "Количество"], [["a", True]], {"copies_column": "Количество"}, "COPIES_INVALID", "B2"),
        (["ID", "Количество"], [["a", 1.5]], {"copies_column": "Количество"}, "COPIES_INVALID", "B2"),
        (["ID", "Количество"], [["a", -1]], {"copies_column": "Количество"}, "COPIES_INVALID", "B2"),
        (["ID", "Количество"], [["a", None]], {"copies_column": "Количество"}, "COPIES_INVALID", "B2"),
    ],
)
def test_reader_reports_strict_header_id_and_copies_errors(
    tmp_path: Path,
    headers: list[object],
    rows: list[list[object]],
    kwargs: dict[str, str],
    code: str,
    cell: str,
) -> None:
    path = tmp_path / f"{code}.xlsx"
    _write_table(path, [headers, *rows])
    with pytest.raises(ProjectError) as caught:
        XlsxReader().read(path, "Лист", **kwargs)
    assert caught.value.diagnostic.code == code
    assert caught.value.diagnostic.cell == cell


def test_reader_rejects_merged_data_ranges(tmp_path: Path) -> None:
    path = tmp_path / "merged.xlsx"
    _write_table(path, [["ID", "Текст"], ["x", "y"]])
    book = load_workbook(path)
    book.active.merge_cells("A2:B2")
    book.save(path)
    book.close()
    with pytest.raises(ProjectError) as caught:
        XlsxReader().read(path, "Лист")
    assert caught.value.diagnostic.code == "MERGED_CELLS_UNSUPPORTED"
    assert caught.value.diagnostic.sheet == "Лист"


def test_real_formula_xml_distinguishes_cached_empty_missing_and_error(tmp_path: Path) -> None:
    path = tmp_path / "formulas.xlsx"
    _write_table(path, [
        ["ID", "Кэш", "Пусто", "НетКэша", "Ошибка"],
        ["x", "=3+4", '=""', "=1+1", "=1/0"],
    ])
    _patch_formula_cells(path)

    row = XlsxReader().read(path, "Лист", id_column="ID").rows[0]
    assert row.values["Кэш"].formula_state is FormulaState.CACHED
    assert value_to_text(row.values["Кэш"].value) == "7"
    assert row.values["Пусто"].formula_state is FormulaState.EMPTY
    assert value_to_text(row.values["Пусто"].value) == ""
    assert row.values["НетКэша"].formula_state is FormulaState.MISSING
    assert row.values["Ошибка"].formula_state is FormulaState.ERROR

    project = _project()
    resolver = BindingResolver(tmp_path, project)
    snapshot = DataSheetSnapshot("main", path, "Лист", "hash", 1, tuple(row.values), (row,))
    empty = resolver.resolve_component(_html_component("<p>{Пусто}</p>"), snapshot, row)
    assert empty.component.elements[0].html == "<p></p>"
    for column, code in (("НетКэша", "FORMULA_RESULT_UNAVAILABLE"), ("Ошибка", "FORMULA_ERROR")):
        with pytest.raises(ProjectError) as caught:
            resolver.resolve_component(_html_component(f"<p>{{{column}}}</p>"), snapshot, row)
        assert caught.value.diagnostic.code == code
        assert caught.value.diagnostic.cell == row.values[column].coordinate
        assert caught.value.diagnostic.node_id == "text"


def test_spreadsheet_text_is_not_reinterpreted_as_template_during_render(tmp_path: Path, monkeypatch) -> None:
    """Excel text is data: braces and Jinja-looking text must remain literal."""
    project = _project(variables={"secret": "НЕ ДОЛЖНО ПОЯВИТЬСЯ"})
    component = _html_component("<p>{Текст}</p>")
    data, row = _data_row(Текст='{{ vars["secret"] }} and {ДругойСтолбец}')
    resolved = BindingResolver(tmp_path, project).resolve_component(component, data, row)
    document = DocumentSnapshot(tmp_path / "components/card.yaml", "", "", component)
    snapshot = ProjectSnapshot(tmp_path, "", "", project, {"card": document})
    captured: list[str] = []

    class _Prepared:
        def draw(self, painter) -> None:
            del painter

    def capture(node, html, loader, owner):
        del node, loader, owner
        captured.append(html)
        return _Prepared()

    monkeypatch.setattr("componentpress.rendering.painter.prepare_html", capture)
    from componentpress.rendering.exporter import render_component_image

    render_component_image(snapshot, "card", component=resolved.component, dpi=96)
    assert captured == ['<p>{{ vars["secret"] }} and {ДругойСтолбец}</p>']


def test_missing_dynamic_image_reports_originating_excel_cell_and_node(tmp_path: Path) -> None:
    project = _project()
    component = ComponentDefinition(
        schema_version=4,
        id="card",
        name="Карта",
        size_mm=SizeMM(width=40.0, height=30.0),
        data=DataBinding(source="main", sheet="Лист"),
        elements=(ImageNode(
            id="art", name="Рисунок", type="image", x_mm=0.0, y_mm=0.0,
            width_mm=20.0, height_mm=20.0, source="assets/images/{Путь}",
        ),),
    )
    data, row = _data_row(Путь="missing.png")
    with pytest.raises(ProjectError) as caught:
        BindingResolver(tmp_path, project).resolve_component(component, data, row)
    diagnostic = caught.value.diagnostic
    assert diagnostic.code == "RESOURCE_MISSING"
    assert (diagnostic.source, diagnostic.sheet, diagnostic.cell, diagnostic.node_id) == (
        "main", "Лист", "A2", "art",
    )


def test_demo_two_sheets_icons_nested_paths_and_template_immutability(tmp_path: Path) -> None:
    root = tmp_path / "demo"
    shutil.copytree(DEMO, root)
    snapshot = project_repository().open(root)
    before_yaml = {name: doc.path.read_bytes() for name, doc in snapshot.documents.items()}
    before_models = {name: doc.model for name, doc in snapshot.documents.items()}
    service = PreviewService(XlsxReader())

    forest = service.select_row(snapshot, "forest-card", row_number=2)
    event = service.select_row(snapshot, "event-card", row_number=2)
    forest_dump = repr(forest.component)
    event_dump = repr(event.component)
    assert "assets/images/animals/enemy/leaf.png" in forest_dump
    assert "Волк" in forest_dump and "Лесной дождь" in event_dump
    assert len(forest.row.instance_id) == 36
    assert forest.row is not None and forest.row.copies == 2
    assert event.row is not None and event.row.copies == 1
    assert {name: doc.path.read_bytes() for name, doc in snapshot.documents.items()} == before_yaml
    assert {name: doc.model for name, doc in snapshot.documents.items()} == before_models


def test_preview_rejects_result_that_becomes_stale_while_reader_is_running(tmp_path: Path) -> None:
    root = tmp_path / "project"
    (root / "components").mkdir(parents=True)
    (root / "data").mkdir()
    component = _html_component("<p>{Текст}</p>")
    project = _project()
    document = DocumentSnapshot(root / "components/card.yaml", "", "", component)
    snapshot = ProjectSnapshot(root, "", "", project, {"card": document})
    data, _row = _data_row(Текст="A")

    class SupersedingReader:
        def __init__(self) -> None:
            self.service: PreviewService | None = None

        def read(self, path, sheet, **kwargs):
            del path, sheet
            assert self.service is not None
            self.service.begin_refresh(root, "card")
            return DataSheetSnapshot(
                data.source, data.path, data.sheet, data.content_hash,
                kwargs["version"], data.headers, data.rows, kwargs["request_id"],
            )

    reader = SupersedingReader()
    service = PreviewService(reader)
    reader.service = service
    request = service.begin_refresh(root, "card")
    with pytest.raises(ProjectError) as caught:
        service.refresh(snapshot, "card", request_id=request)
    assert caught.value.diagnostic.code == "DATA_REFRESH_STALE"


def test_cli_preview_export_and_gui_canvas_use_same_resolved_pixels(tmp_path: Path, qtbot) -> None:
    root = tmp_path / "demo"
    shutil.copytree(DEMO, root)
    snapshot = project_repository().open(root)
    service = PreviewService(XlsxReader())
    expected = tmp_path / "service.png"
    cli = tmp_path / "cli.png"
    service.export_png(snapshot, "forest-card", expected, row_number=2, dpi=96)
    assert cli_main([
        "render", str(root), "--component", "forest-card", "--row", "2",
        "--dpi", "96", "--output", str(cli),
    ]) == 0
    assert cli.read_bytes() == expected.read_bytes()

    window = MainWindow(project_controller())
    qtbot.addWidget(window)
    window.load_session(project_service().open_session(root))
    tab = window.open_component("forest-card")
    assert tab is not None
    fox_index = window.data_row.findData(3)
    window.data_row.setCurrentIndex(fox_index)
    canvas_image = tab.canvas._pixmap_item.pixmap().toImage()
    exported_image = QImage(str(expected))
    assert canvas_image.size() == exported_image.size()
    assert canvas_image.convertToFormat(exported_image.format()) == exported_image
    window._clear_tabs(discard=True)


def test_static_templates_still_resolve_and_invalid_syntax_is_a_project_error(tmp_path: Path, monkeypatch) -> None:
    project = _project(variables={"edition": "<b>Статическое значение</b>"})
    component = _html_component('<p>{{ vars["edition"] }}</p>').model_copy(update={"data": None})
    document = DocumentSnapshot(tmp_path / "components/card.yaml", "", "", component)
    snapshot = ProjectSnapshot(tmp_path, "", "", project, {"card": document})
    captured: list[str] = []

    class _Prepared:
        def draw(self, painter) -> None:
            del painter

    def capture(node, html, loader, owner):
        del node, loader, owner
        captured.append(html)
        return _Prepared()

    monkeypatch.setattr("componentpress.rendering.painter.prepare_html", capture)
    from componentpress.rendering.exporter import render_component_image

    render_component_image(snapshot, "card", dpi=96)
    assert captured == ["<p>&lt;b&gt;Статическое значение&lt;/b&gt;</p>"]

    invalid = component.model_copy(update={
        "elements": (component.elements[0].model_copy(update={"html": "<p>{{ vars.edition }}</p>"}),),
    })
    invalid_snapshot = ProjectSnapshot(
        tmp_path, "", "", project,
        {"card": DocumentSnapshot(document.path, "", "", invalid)},
    )
    with pytest.raises(ProjectError) as caught:
        render_component_image(invalid_snapshot, "card", dpi=96)
    assert caught.value.diagnostic.code == "TEMPLATE_SYNTAX"


@pytest.mark.parametrize("kind", ["image", "html", "icon"])
@pytest.mark.parametrize(
    ("relative", "expected"),
    [("assets/images/missing.png", "RESOURCE_MISSING"), ("assets/images/broken.png", "RESOURCE_INVALID")],
)
def test_dynamic_resource_errors_keep_excel_context_for_all_render_contexts(
    tmp_path: Path,
    kind: str,
    relative: str,
    expected: str,
) -> None:
    broken = tmp_path / "assets" / "images" / "broken.png"
    broken.parent.mkdir(parents=True, exist_ok=True)
    broken.write_bytes(b"not an image")
    project = _project()
    if kind == "image":
        component = ComponentDefinition(
            schema_version=4, id="card", name="Карта", size_mm=SizeMM(width=40.0, height=30.0),
            data=DataBinding(source="main", sheet="Лист"),
            elements=(ImageNode(
                id="subject", name="Ресурс", type="image", x_mm=0.0, y_mm=0.0,
                width_mm=20.0, height_mm=20.0, source="{Значение}",
            ),),
        )
        value = relative
    elif kind == "html":
        component = _html_component('<p><img src="{Значение}" /></p>')
        value = relative
    else:
        label = "ic_test"
        project = project.model_copy(update={
            "icons": {label: IconDefinition(path=relative, width_mm=3.0, height_mm=3.0)},
        })
        component = _html_component("<p>{Значение}</p>")
        value = label
    data, row = _data_row(Значение=value)
    with pytest.raises(ProjectError) as caught:
        BindingResolver(tmp_path, project).resolve_component(component, data, row)
    diagnostic = caught.value.diagnostic
    assert diagnostic.code == expected
    assert (diagnostic.source, diagnostic.sheet, diagnostic.cell, diagnostic.node_id) == (
        "main", "Лист", "A2", "subject" if kind == "image" else "text",
    )


def test_static_gui_canvas_matches_static_cli_export_with_project_variables(tmp_path: Path, qtbot) -> None:
    root = tmp_path / "static-demo"
    shutil.copytree(STATIC_DEMO, root)
    snapshot = project_repository().open(root)
    exported = tmp_path / "static.png"
    from componentpress.rendering.exporter import render_component

    render_component(snapshot, "forest-card", exported, dpi=96)
    window = MainWindow(project_controller())
    qtbot.addWidget(window)
    window.load_session(project_service().open_session(root))
    tab = window.open_component("forest-card")
    assert tab is not None
    canvas_image = tab.canvas._pixmap_item.pixmap().toImage()
    exported_image = QImage(str(exported))
    assert canvas_image.size() == exported_image.size()
    assert canvas_image.convertToFormat(exported_image.format()) == exported_image
    window._clear_tabs(discard=True)


def test_preview_cache_uses_project_source_and_component_sheet(tmp_path: Path) -> None:
    root = tmp_path / "cache-project"
    (root / "components").mkdir(parents=True)
    (root / "data").mkdir()
    (root / "data/a.xlsx").write_bytes(b"a")
    (root / "data/b.xlsx").write_bytes(b"b")
    component = ComponentDefinition(
        schema_version=4, id="card", name="Карта", size_mm=SizeMM(width=40, height=30),
        data=DataBinding(sheet="Лист"),
        elements=(HtmlNode(id="text", name="Текст", type="html", x_mm=0, y_mm=0,
            width_mm=40, height_mm=30, font_family="Arial", font_size_pt=8,
            html="", content_mode="column", content_column="Значение"),),
    )
    project = _project().model_copy(update={"data_source": DataSource(path="data/a.xlsx")})
    document = DocumentSnapshot(root / "components/card.yaml", "", "", component)
    base = ProjectSnapshot(root, "", "", project, {"card": document})

    class RecordingReader:
        def __init__(self) -> None:
            self.calls: list[tuple[str, str]] = []

        def read(self, path, sheet, **kwargs):
            call = (Path(path).name, sheet)
            self.calls.append(call)
            row = DataRow(2, repr(call), 1, {"Значение": CellValue(repr(call), "A2")})
            return DataSheetSnapshot(
                kwargs["source"], Path(path), sheet, repr(call), kwargs["version"],
                ("Значение",), (row,), kwargs["request_id"],
            )

    reader = RecordingReader()
    service = PreviewService(reader)
    cases: list[tuple[ProjectSnapshot, ComponentDefinition]] = [(base, component)]
    cases.append((base, component.model_copy(update={"data": component.data.model_copy(update={"sheet": "Другой"})})))
    path_project = project.model_copy(update={"data_source": DataSource(path="data/b.xlsx")})
    cases.append((ProjectSnapshot(root, "", "", path_project, {"card": document}), component))

    snapshots = [service.data(snapshot, "card", component=model) for snapshot, model in cases]
    assert len(reader.calls) == len(cases)
    assert len({item.rows[0].instance_id for item in snapshots if item is not None}) == 2
    for snapshot, model in cases:
        service.data(snapshot, "card", component=model)
    assert len(reader.calls) == len(cases), "повтор того же identity должен использовать свой кэш"


def test_gui_prod_test_preview_filters_validated_snapshot(tmp_path: Path, qtbot) -> None:
    root = tmp_path / "gui-invalidation"
    shutil.copytree(DEMO, root)
    controller = project_controller()
    window = MainWindow(controller)
    qtbot.addWidget(window)
    window.load_session(project_service().open_session(root))
    try:
        tab = window.open_component("forest-card")
        assert tab is not None and tab.preview_row_number == 2
        data = controller.preview.data(window.session.snapshot, "forest-card")
        assert data is not None and len(data.rows) == 2
        window.data_mode.setCurrentIndex(1)
        assert window.data_row.count() == 2
        window.data_row.setCurrentIndex(1)
        assert tab.preview_row_number == 3
        assert "Лиса" in repr(tab.resolved_component)
        assert controller.preview.data(window.session.snapshot, "forest-card") is data
    finally:
        window._clear_tabs(discard=True)


@pytest.mark.parametrize(
    ("target", "expected_cell"),
    [("html", "B2"), ("id", "A2"), ("copies", "C2")],
)
def test_ordinary_excel_error_cells_block_every_consumer_with_exact_address(
    tmp_path: Path,
    target: str,
    expected_cell: str,
) -> None:
    path = tmp_path / f"ordinary-error-{target}.xlsx"
    values: list[object] = ["card", "normal", 1]
    values[{"id": 0, "html": 1, "copies": 2}[target]] = "#REF!"
    _write_table(path, [["ID", "Текст", "Количество"], values])
    reader = XlsxReader()
    kwargs = {
        "source": "main",
        "id_column": "ID" if target == "id" else None,
        "copies_column": "Количество" if target == "copies" else None,
    }
    if target in {"id", "copies"}:
        with pytest.raises(ProjectError) as caught:
            reader.read(path, "Лист", **kwargs)
    else:
        data = reader.read(path, "Лист", **kwargs)
        assert data.rows[0].values["Текст"].formula is None
        assert data.rows[0].values["Текст"].formula_state is FormulaState.ERROR
        with pytest.raises(ProjectError) as caught:
            BindingResolver(tmp_path, _project()).resolve_component(
                _html_component("<p>{Текст}</p>"), data, data.rows[0],
            )
    diagnostic = caught.value.diagnostic
    assert diagnostic.code == "FORMULA_ERROR"
    assert diagnostic.cell == expected_cell
    assert diagnostic.sheet == "Лист"


@pytest.mark.parametrize(
    ("node", "expected_field", "expected_node"),
    [
        (
            ImageNode(
                id="art", name="Рисунок", type="image", x_mm=0.0, y_mm=0.0,
                width_mm=20.0, height_mm=20.0, source="assets/images/{Путь}",
            ),
            "elements.art.source",
            "art",
        ),
        (
            HtmlNode(
                id="body", name="Текст", type="html", x_mm=0.0, y_mm=0.0,
                width_mm=30.0, height_mm=20.0, font_family="Arial", font_size_pt=8.0,
                html='<p><img src="assets/images/{Путь}" /></p>',
            ),
            "elements.body.html.img.src",
            "body",
        ),
    ],
)
def test_path_invalid_from_excel_keeps_full_origin(
    tmp_path: Path,
    node,
    expected_field: str,
    expected_node: str,
) -> None:
    component = ComponentDefinition(
        schema_version=4, id="card", name="Карта", size_mm=SizeMM(width=40.0, height=30.0),
        data=DataBinding(source="main", sheet="Лист"), elements=(node,),
    )
    data, row = _data_row(Путь="../outside.png")
    with pytest.raises(ProjectError) as caught:
        BindingResolver(tmp_path, _project()).resolve_component(component, data, row)
    diagnostic = caught.value.diagnostic
    assert diagnostic.code == "PATH_INVALID"
    assert diagnostic.field == expected_field
    assert (diagnostic.source, diagnostic.sheet, diagnostic.cell, diagnostic.node_id) == (
        "main", "Лист", "A2", expected_node,
    )


def test_escape_markers_apply_only_to_template_not_inserted_excel_or_variable_bytes(tmp_path: Path) -> None:
    excel_value = r"excel \{one\} \}two\{"
    variable_value = r"variable \{three\} \}four\{"
    project = _project(variables={"literal": variable_value})
    component = _html_component(
        r'<p>template \{five\} \}six\{ | {Текст} | {{ vars["literal"] }}</p>'
    )
    data, row = _data_row(Текст=excel_value)
    resolved = BindingResolver(tmp_path, project).resolve_component(component, data, row)
    assert resolved.component.elements[0].html == (
        r"<p>template {five} }six{ | excel \{one\} \}two\{ | variable \{three\} \}four\{</p>"
    )

    assert resolve_static(
        r'template \{five\} | {{ vars["literal"] }}',
        {"literal": variable_value},
        html=False,
        owner=tmp_path / "component.yaml",
        field="html",
    ) == r"template {five} | variable \{three\} \}four\{"
