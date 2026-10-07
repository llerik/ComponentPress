"""Acceptance checks for Prod/Test builds and PNG ZIP export."""

from __future__ import annotations

import json
from pathlib import Path
import shutil
import zipfile

from openpyxl import load_workbook
import pytest

from componentpress.application.build_service import BuildService
from componentpress.application.contracts import BuildRequest
from componentpress.application.png_archive import export_png_archive
from componentpress.application.preview_service import PreviewService
from componentpress.bootstrap import project_repository, project_service
from componentpress.data_sources import XlsxReader
from componentpress.domain.diagnostics import ProjectError
from componentpress.editor.controllers import ProjectController
from componentpress.editor.main_window import MainWindow
from componentpress.execution.cancellation import CancellationToken
from componentpress.rendering.exporter import ensure_gui_application


DEMO = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"
STATIC_DEMO = Path(__file__).resolve().parents[1] / "examples" / "demo-game"


def _copy_demo(tmp_path: Path) -> Path:
    root = tmp_path / "Игра режима"
    shutil.copytree(DEMO, root, ignore=shutil.ignore_patterns("output", "archive", ".componentpress"))
    return root


def test_test_build_selects_test_rows_reports_mode_and_zip_manifest(tmp_path: Path) -> None:
    ensure_gui_application()
    root = _copy_demo(tmp_path)
    service = project_service()
    session = service.open_session(root)
    result = BuildService(service).run(session, BuildRequest(mode="test", component_ids=("forest-card",)))
    assert result.status == "succeeded", result.diagnostics
    assert result.mode == "test"
    report = json.loads(result.report_path.read_text(encoding="utf-8"))
    assert report["mode"] == "test"
    assert report["report_version"] == 2
    assert report["counts"]["png"] == 2
    assert [(item["row_number"], item["copies"]) for item in report["instances"]] == [(2, 1), (3, 3)]
    assert report["counts"]["copies"] == 4

    archive_path = tmp_path / "Экспорт" / "карты-test.zip"
    archived = export_png_archive(session.snapshot, archive_path, mode="test", component_ids=("forest-card",))
    assert (archived.mode, archived.png_count, archived.copies) == ("test", 2, 4)
    with zipfile.ZipFile(archive_path) as archive:
        assert archive.testzip() is None
        assert archive.namelist() == [
            "images/0001-forest-card-" + report["instances"][0]["instance_id"] + ".png",
            "images/0002-forest-card-" + report["instances"][1]["instance_id"] + ".png",
            "manifest.json",
        ]
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["mode"] == "test"
        assert manifest["counts"] == {"png": 2, "copies": 4}
        assert all(record["sha256"] for record in manifest["instances"])


def test_both_modes_validate_resources_in_zero_quantity_rows(tmp_path: Path) -> None:
    ensure_gui_application()
    root = _copy_demo(tmp_path)
    workbook_path = root / "data" / "game.xlsx"
    workbook = load_workbook(workbook_path)
    workbook["Карты"]["D3"] = "assets/images/missing-zero-row.png"
    workbook.save(workbook_path)
    workbook.close()
    service = project_service()
    session = service.open_session(root)
    builder = BuildService(service)
    for mode in ("prod", "test"):
        result = builder.run(session, BuildRequest(mode=mode, component_ids=("forest-card",)))
        assert result.status == "failed"
        assert any(item.code == "RESOURCE_MISSING" for item in result.diagnostics)


def test_invalid_mode_and_png_archive_conflict_are_rejected(tmp_path: Path) -> None:
    ensure_gui_application()
    root = _copy_demo(tmp_path)
    service = project_service()
    session = service.open_session(root)
    with pytest.raises(ProjectError) as error:
        BuildService(service).start(session, BuildRequest(mode="debug"))
    assert error.value.diagnostic.code == "COPY_MODE_INVALID"

    target = tmp_path / "existing.zip"
    target.write_bytes(b"keep")
    with pytest.raises(ProjectError) as error:
        export_png_archive(session.snapshot, target, mode="prod", component_ids=("forest-card",))
    assert error.value.diagnostic.code == "ARCHIVE_EXISTS"
    assert target.read_bytes() == b"keep"


def test_empty_selected_mode_and_cancelled_archive_publish_nothing(tmp_path: Path) -> None:
    ensure_gui_application()
    root = _copy_demo(tmp_path)
    workbook_path = root / "data" / "game.xlsx"
    workbook = load_workbook(workbook_path)
    workbook["Карты"]["J2"] = 0
    workbook["Карты"]["J3"] = 0
    workbook.save(workbook_path)
    workbook.close()
    service = project_service()
    session = service.open_session(root)
    result = BuildService(service).run(session, BuildRequest(mode="test", component_ids=("forest-card",)))
    assert result.status == "failed"
    assert any(item.code == "BUILD_EMPTY" for item in result.diagnostics)

    token = CancellationToken()
    token.cancel()
    target = tmp_path / "cancelled.zip"
    with pytest.raises(ProjectError) as error:
        export_png_archive(session.snapshot, target, mode="prod", component_ids=("forest-card",), cancellation=token)
    assert error.value.diagnostic.code == "BUILD_CANCELLED"
    assert not target.exists()


def test_identical_quantity_column_names_apply_to_both_modes(tmp_path: Path) -> None:
    ensure_gui_application()
    root = _copy_demo(tmp_path)
    project_path = root / "project.yaml"
    project_text = project_path.read_text(encoding="utf-8").replace("test: Debug", "test: Prod")
    project_path.write_text(project_text, encoding="utf-8")
    service = project_service()
    session = service.open_session(root)
    result = BuildService(service).run(session, BuildRequest(mode="test", component_ids=("forest-card",)))
    assert result.status == "succeeded", result.diagnostics
    report = json.loads(result.report_path.read_text(encoding="utf-8"))
    assert report["mode"] == "test"
    assert report["counts"]["copies"] == 2
    assert [(item["row_number"], item["copies"]) for item in report["instances"]] == [(2, 2)]


def test_static_component_is_one_instance_in_test_mode(tmp_path: Path) -> None:
    ensure_gui_application()
    root = tmp_path / "Статический проект"
    shutil.copytree(STATIC_DEMO, root)
    service = project_service()
    session = service.open_session(root)
    component_id = next(iter(session.documents))
    result = BuildService(service).run(session, BuildRequest(mode="test", component_ids=(component_id,)))
    assert result.status == "succeeded", result.diagnostics
    report = json.loads(result.report_path.read_text(encoding="utf-8"))
    assert report["mode"] == "test"
    assert report["counts"]["instances"] == report["counts"]["copies"] == 1


def test_cli_build_and_render_accept_test_mode(tmp_path: Path) -> None:
    ensure_gui_application()
    from componentpress.cli import main

    root = _copy_demo(tmp_path)
    assert main(["build", str(root), "--component", "forest-card", "--mode", "test", "--workers", "1"]) == 0
    reports = list((root / "output").glob("build-*/build-report.json"))
    assert len(reports) == 1
    report = json.loads(reports[0].read_text(encoding="utf-8"))
    assert report["mode"] == "test"
    rendered = tmp_path / "test-row.png"
    assert main(["render", str(root), "--component", "forest-card", "--row", "3", "--mode", "test", "--output", str(rendered)]) == 0
    assert rendered.is_file()


def test_cached_preview_lookup_does_not_read_excel(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from componentpress.application.preview_service import PreviewService
    from componentpress.data_sources import XlsxReader

    snapshot = project_repository().open(_copy_demo(tmp_path))
    preview = PreviewService(XlsxReader())
    cached = preview.refresh(snapshot, "forest-card")
    assert cached is not None

    def fail_read(*_args, **_kwargs):
        raise AssertionError("cached mode switch attempted an XLSX read")

    monkeypatch.setattr(preview.reader, "read", fail_read)
    assert preview.cached_data(snapshot, "forest-card") is cached


def test_gui_mode_switch_uses_cached_excel_snapshot(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, qtbot) -> None:
    ensure_gui_application()
    service = project_service()
    session = service.open_session(_copy_demo(tmp_path))
    window = MainWindow(ProjectController(service, PreviewService(XlsxReader()), BuildService(service)))
    qtbot.addWidget(window)
    window.load_session(session)
    tab = window.open_component("forest-card")
    assert tab is not None and window.controller.preview is not None

    def fail_read(*_args, **_kwargs):
        raise AssertionError("GUI mode switch attempted an XLSX read")

    monkeypatch.setattr(window.controller.preview.reader, "read", fail_read)
    window.data_mode.setCurrentIndex(1)
    assert window.data_mode.currentData() == "test"
    assert tab.preview_row_number == 2
