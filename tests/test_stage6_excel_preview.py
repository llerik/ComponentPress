"""Coder checks for stage six; independent acceptance remains with the tester role."""

from __future__ import annotations

from datetime import date
from pathlib import Path
import re
import shutil
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from componentpress.application.preview_service import PreviewService
from componentpress.bindings.resolver import BindingResolver
from componentpress.bootstrap import project_repository
from componentpress.bootstrap import project_controller, project_service
from componentpress.data_sources import XlsxReader
from componentpress.domain.diagnostics import ProjectError
from componentpress.domain.values import CellValue, DataRow, DataSheetSnapshot, FormulaState, value_to_text
from componentpress.domain.component import ComponentDefinition, DataBinding, SizeMM
from componentpress.domain.nodes import HtmlNode, ImageNode
from componentpress.domain.project import IconDefinition, ProjectDefinition
from componentpress.editor.main_window import MainWindow
from PySide6.QtGui import QFont, QTextCursor
from PySide6.QtWidgets import QFileDialog


EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "demo-game"
EXCEL_EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "demo-excel-game"


def _patch_formula_cache(path: Path, coordinate: str, value: str | None, *, error: bool = False, remove_value: bool = False) -> None:
    temporary = path.with_suffix(".patched.xlsx")
    with ZipFile(path) as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                pattern = rf'(<c[^>]*r="{coordinate}"[^>]*>.*?<f[^>]*>.*?</f>)(?:<v(?:>.*?</v>|\s*/>))?(</c>)'
                cached = "" if remove_value else f"<v>{'' if value is None else value}</v>"
                replacement = rf"\1{cached}\2"
                text, count = re.subn(pattern, replacement, text, count=1)
                assert count == 1
                if error:
                    text = re.sub(rf'<c([^>]*r="{coordinate}"[^>]*)>', lambda match: '<c' + re.sub(r'\s+t="[^"]*"', '', match.group(1)) + ' t="e">', text, count=1)
                data = text.encode("utf-8")
            target.writestr(item, data)
    temporary.replace(path)


def _book(path: Path, *, formula_cache: str | None = "7", missing: bool = False, error: bool = False) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Карты"
    sheet.append(["ID", "Название", "Описание", "Изображение", "Сила", "Количество", "Дата", "Флаг", "Prod", "Debug", "Название HTML", "Описание HTML"])
    sheet.append(["wolf", "Волк", r"Атака ic_leaf_6 и ic_leaf_3; буквально \ic_leaf_3", "assets/images/animals/enemy/leaf.png", "=3+4", 2, date(2026, 10, 3), True, 2, 1, "<p><b>Волк</b></p>", "<p>Атака ic_leaf_6 и ic_leaf_3; буквально \\ic_leaf_3</p>"])
    sheet.append(["fox", "Лиса", "Тихая", "assets/images/animals/enemy/leaf.png", 2.5, 0, None, False, 0, 3, "<p><b>Лиса</b></p>", "<p>Тихая</p>"])
    events = book.create_sheet("События")
    events.append(["ID", "Название", "Prod", "Debug", "Название HTML", "Описание HTML"])
    events.append(["event-1", "Дождь", 1, 0, "<p>Дождь</p>", "<p>Событие</p>"])
    book.save(path)
    book.close()
    _patch_formula_cache(path, "E2", "#VALUE!" if error else formula_cache, error=error, remove_value=missing)


def _bound_project(tmp_path: Path) -> Path:
    root = tmp_path / "Проект Excel"
    shutil.copytree(EXAMPLE, root)
    (root / "data").mkdir(exist_ok=True)
    _book(root / "data/game.xlsx")
    from io import StringIO
    from ruamel.yaml import YAML
    yaml = YAML(typ="rt")
    project = root / "project.yaml"
    tree = yaml.load(project.read_text(encoding="utf-8"))
    tree["data_source"] = {"path": "data/game.xlsx"}
    tree["icons"]["ic_leaf_3"] = {"path": "assets/images/animals/enemy/leaf.png", "width_mm": 3, "height_mm": 3}
    output = StringIO()
    yaml.dump(tree, output)
    project.write_text(output.getvalue(), encoding="utf-8")
    component = root / "components/forest-card.yaml"
    tree = yaml.load(component.read_text(encoding="utf-8"))
    tree["data"] = {"sheet": "Карты"}
    tree["elements"][0]["source"] = {"mode": "column", "column": "Изображение"}
    tree["elements"][1]["children"][0]["content"] = {"mode": "column", "column": "Название HTML"}
    tree["elements"][1]["children"][1]["content"] = {"mode": "column", "column": "Описание HTML"}
    output = StringIO()
    yaml.dump(tree, output)
    component.write_text(output.getvalue(), encoding="utf-8")
    return root


def test_reader_cached_formula_types_copies_and_file_release(tmp_path: Path) -> None:
    path = tmp_path / "данные.xlsx"
    _book(path)
    snapshot = XlsxReader().read(path, "Карты", source="main", id_column="ID", copies_column="Количество")
    first = snapshot.rows[0]
    assert snapshot.headers[:3] == ("ID", "Название", "Описание")
    assert first.instance_id == "wolf" and first.copies == 2
    assert first.values["Сила"].formula_state is FormulaState.CACHED
    assert value_to_text(first.values["Сила"].value) == "7"
    assert value_to_text(first.values["Дата"].value) == "2026-10-03"
    assert value_to_text(first.values["Флаг"].value) == "true"
    assert value_to_text(snapshot.rows[1].values["Сила"].value) == "2.5"
    moved = path.with_name("после-чтения.xlsx")
    path.replace(moved)
    assert moved.is_file()


def test_formula_missing_empty_and_error_are_distinct(tmp_path: Path) -> None:
    missing = tmp_path / "missing.xlsx"
    empty = tmp_path / "empty.xlsx"
    error = tmp_path / "error.xlsx"
    _book(missing, missing=True)
    _book(empty, formula_cache=None)
    _book(error, error=True)
    reader = XlsxReader()
    assert reader.read(missing, "Карты").rows[0].values["Сила"].formula_state is FormulaState.MISSING
    assert reader.read(empty, "Карты").rows[0].values["Сила"].formula_state is FormulaState.EMPTY
    assert reader.read(error, "Карты").rows[0].values["Сила"].formula_state is FormulaState.ERROR


def test_reader_rejects_headers_ids_and_copies_with_exact_cells(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Карты"
    sheet.append(["ID", "Название", "Название", "Количество"])
    sheet.append(["same", "A", "B", -1])
    book.save(path)
    book.close()
    try:
        XlsxReader().read(path, "Карты", id_column="ID", copies_column="Количество")
    except ProjectError as exc:
        assert exc.diagnostic.code == "HEADER_DUPLICATE"
        assert exc.diagnostic.cell == "C1"
    else:
        raise AssertionError("повторный заголовок должен быть отклонён")


def test_resolver_icons_tables_dynamic_paths_and_immutable_template(tmp_path: Path) -> None:
    root = _bound_project(tmp_path)
    snapshot = project_repository().open(root)
    original = snapshot.documents["forest-card"].model
    data = XlsxReader().read(root / "data/game.xlsx", "Карты", source="main", id_column="ID", copies_column="Количество")
    resolved = BindingResolver(root, snapshot.model).resolve_component(original, data, data.rows[0], owner=snapshot.documents["forest-card"].path)
    dump = repr(resolved.component)
    assert "Волк" in dump and 'alt="ic_leaf_6"' in dump and ">7<" not in dump
    heading = resolved.component.elements[1].children[0]
    assert heading.html == "<p><b>Волк</b></p>"
    from componentpress.rendering.exporter import ensure_gui_application
    from componentpress.rendering.html_document import prepare_html
    from componentpress.rendering.resources import ProjectResourceLoader
    ensure_gui_application()
    rendered = prepare_html(
        heading.model_copy(update={"height_mm": 20}), heading.html,
        ProjectResourceLoader(root), snapshot.documents["forest-card"].path,
    ).document
    assert rendered.toPlainText().strip() == "Волк"
    cursor = QTextCursor(rendered)
    cursor.movePosition(QTextCursor.MoveOperation.NextCharacter, QTextCursor.MoveMode.KeepAnchor)
    assert cursor.charFormat().fontWeight() >= QFont.Weight.Bold
    assert 'width=&quot;' not in dump
    assert "ic_leaf_3" in dump  # escaped label remains literal
    assert "assets/images/animals/enemy/leaf.png" in dump
    assert resolved.component != original
    assert snapshot.documents["forest-card"].model == original


def test_unknown_icon_reports_sheet_row_column_and_node(tmp_path: Path) -> None:
    root = _bound_project(tmp_path)
    path = root / "data/game.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Карты"
    sheet.append(["ID", "Название", "Описание", "Изображение", "Сила", "Количество", "Prod", "Debug", "Название HTML", "Описание HTML"])
    sheet.append(["x", "X", "bad ic_unknown", "assets/images/animals/enemy/leaf.png", 1, 1, 1, 1, "<p>X</p>", "<p>bad ic_unknown</p>"])
    book.save(path)
    book.close()
    snapshot = project_repository().open(root)
    data = XlsxReader().read(path, "Карты", source="main", id_column="ID", copies_column="Количество")
    try:
        BindingResolver(root, snapshot.model).resolve_component(snapshot.documents["forest-card"].model, data, data.rows[0], owner=snapshot.documents["forest-card"].path)
    except ProjectError as exc:
        diagnostic = exc.diagnostic
        assert diagnostic.code == "ICON_UNKNOWN"
        assert (diagnostic.sheet, diagnostic.cell, diagnostic.node_id) == ("Карты", "J2", "caption")
    else:
        raise AssertionError("неизвестная иконка должна быть ошибкой")


def test_preview_select_refresh_stale_protection_and_png(tmp_path: Path) -> None:
    root = _bound_project(tmp_path)
    snapshot = project_repository().open(root)
    service = PreviewService(XlsxReader())
    first = service.select_row(snapshot, "forest-card", row_number=2)
    second = service.select_row(snapshot, "forest-card", row_number=3, mode="test")
    assert first.row is not None and second.row is not None
    assert first.row.copies == 2 and second.row.test_copies == 3
    assert first.component != second.component
    first_png = tmp_path / "wolf.png"
    second_png = tmp_path / "fox.png"
    service.export_png(snapshot, "forest-card", first_png, row_number=2, dpi=96)
    service.export_png(snapshot, "forest-card", second_png, row_number=3, mode="test", dpi=96)
    assert first_png.read_bytes() != second_png.read_bytes()
    old = service.begin_refresh(root, "forest-card")
    service.begin_refresh(root, "forest-card")
    try:
        service.refresh(snapshot, "forest-card", request_id=old)
    except ProjectError as exc:
        assert exc.diagnostic.code == "DATA_REFRESH_STALE"
    else:
        raise AssertionError("устаревшее обновление должно быть отброшено")


def test_gui_data_panel_switch_refresh_insert_and_export(tmp_path: Path, qtbot, monkeypatch) -> None:
    root = tmp_path / "GUI Excel"
    shutil.copytree(EXCEL_EXAMPLE, root)
    controller = project_controller()
    window = MainWindow(controller)
    qtbot.addWidget(window)
    window.load_session(project_service().open_session(root))
    window.show()
    tab = window.open_component("forest-card")
    assert tab is not None
    assert window.data_source.currentData() == "main"
    assert window.data_sheet.currentData() == "Карты"
    assert window.data_row.count() == 1
    assert tab.preview_row_number == 2

    before_version = controller.preview.data(window.session.snapshot, "forest-card").version
    window.data_mode.setCurrentIndex(1)
    assert window.data_row.count() == 2
    window.data_row.setCurrentIndex(1)
    assert tab.preview_row_number == 3
    assert tab.resolved_component is not None and "Лиса" in repr(tab.resolved_component)
    window._refresh_data(force=True)
    assert controller.preview.data(window.session.snapshot, "forest-card").version == before_version + 1

    index = window.data_column.findData("Название")
    window.data_column.setCurrentIndex(index)
    window._set_selection(("caption",), source="tree")
    window._insert_column_binding()
    bound = window.session.documents["forest-card"].model.elements[1].children[1]
    assert bound.content_mode == "column" and bound.content_column == "Название"
    assert window.data_chain.text() == "main → Карты → Название"

    output = tmp_path / "gui-instance.png"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *args, **kwargs: (str(output), "PNG (*.png)")))
    window._export_instance()
    assert output.is_file()
    window._clear_tabs(discard=True)


def test_dynamic_html_image_missing_keeps_excel_origin(tmp_path: Path) -> None:
    project = ProjectDefinition(schema_version=4, id="game", name="Game", version="0.1.0")
    component = ComponentDefinition(
        schema_version=4, id="card", name="Card", size_mm=SizeMM(width=40, height=30),
        data=DataBinding(source="main", sheet="Лист"),
        elements=(HtmlNode(
            id="body", name="Body", type="html", x_mm=0, y_mm=0,
            width_mm=30, height_mm=20, font_family="Arial", font_size_pt=8,
            html='<p><img src="assets/images/{Путь}" /></p>',
        ),),
    )
    row = DataRow(2, "2", 1, {"Путь": CellValue("missing.png", "A2")})
    data = DataSheetSnapshot("main", tmp_path / "data.xlsx", "Лист", "hash", 1, ("Путь",), (row,))
    try:
        BindingResolver(tmp_path, project).resolve_component(component, data, row)
    except ProjectError as exc:
        diagnostic = exc.diagnostic
        assert diagnostic.code == "RESOURCE_MISSING"
        assert (diagnostic.source, diagnostic.sheet, diagnostic.cell, diagnostic.node_id) == ("main", "Лист", "A2", "body")
    else:
        raise AssertionError("отсутствующий динамический img.src должен быть ошибкой")


def test_missing_and_broken_icon_keep_excel_origin(tmp_path: Path) -> None:
    broken = tmp_path / "assets/images/broken.png"
    broken.parent.mkdir(parents=True)
    broken.write_bytes(b"not a png")
    component = ComponentDefinition(
        schema_version=4, id="card", name="Card", size_mm=SizeMM(width=40, height=30),
        data=DataBinding(source="main", sheet="Лист"),
        elements=(HtmlNode(
            id="body", name="Body", type="html", x_mm=0, y_mm=0,
            width_mm=30, height_mm=20, font_family="Arial", font_size_pt=8, html="<p>{Текст}</p>",
        ),),
    )
    for name, path, expected in (
        ("ic_missing", "assets/images/missing.png", "RESOURCE_MISSING"),
        ("ic_broken", "assets/images/broken.png", "RESOURCE_INVALID"),
    ):
        project = ProjectDefinition(
            schema_version=4, id="game", name="Game", version="0.1.0",
            icons={name: IconDefinition(path=path, width_mm=3, height_mm=3)},
        )
        row = DataRow(2, "2", 1, {"Текст": CellValue(name, "B2")})
        data = DataSheetSnapshot("main", tmp_path / "data.xlsx", "Лист", "hash", 1, ("Текст",), (row,))
        try:
            BindingResolver(tmp_path, project).resolve_component(component, data, row)
        except ProjectError as exc:
            diagnostic = exc.diagnostic
            assert diagnostic.code == expected
            assert (diagnostic.source, diagnostic.sheet, diagnostic.cell, diagnostic.node_id) == ("main", "Лист", "B2", "body")
        else:
            raise AssertionError("невалидный ресурс иконки должен быть ошибкой")


def test_preview_cache_identity_changes_with_unsaved_binding_model(tmp_path: Path) -> None:
    root = tmp_path / "Cache binding"
    shutil.copytree(EXCEL_EXAMPLE, root)
    snapshot = project_repository().open(root)
    service = PreviewService(XlsxReader())
    forest = snapshot.documents["forest-card"].model
    first = service.select_row(snapshot, "forest-card")
    assert first.row is not None and first.row.row_number == 2

    event = snapshot.documents["event-card"].model
    switched = event.model_copy(update={
        "id": "forest-card",
        "data": event.data.model_copy(update={"sheet": "События"}),
    })
    changed_data = service.data(snapshot, "forest-card", component=switched)
    changed = service.select_row(snapshot, "forest-card", component=switched)
    assert changed_data is not None and changed_data.sheet == "События"
    assert changed.row is not None and changed.row.row_number == 2


def test_ordinary_excel_error_blocks_html_and_configured_id(tmp_path: Path) -> None:
    path = tmp_path / "errors.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист"
    sheet.append(["ID", "Текст"])
    sheet.append(["ok", "#VALUE!"])
    book.save(path)
    book.close()
    reader = XlsxReader()
    data = reader.read(path, "Лист", source="main", id_column="ID")
    assert data.rows[0].values["Текст"].formula_state is FormulaState.ERROR
    project = ProjectDefinition(schema_version=4, id="game", name="Game", version="0.1.0")
    component = ComponentDefinition(
        schema_version=4, id="card", name="Card", size_mm=SizeMM(width=40, height=30),
        data=DataBinding(source="main", sheet="Лист", id_column="ID"),
        elements=(HtmlNode(
            id="body", name="Body", type="html", x_mm=0, y_mm=0, width_mm=30,
            height_mm=20, font_family="Arial", font_size_pt=8, html="<p>{Текст}</p>",
        ),),
    )
    try:
        BindingResolver(tmp_path, project).resolve_component(component, data, data.rows[0])
    except ProjectError as exc:
        assert exc.diagnostic.code == "FORMULA_ERROR" and exc.diagnostic.cell == "B2"
    else:
        raise AssertionError("обычная Excel error-cell должна блокировать HTML")

    book = Workbook()
    sheet = book.active
    sheet.title = "Лист"
    sheet.append(["ID", "Текст"])
    sheet.append(["#VALUE!", "ok"])
    book.save(path)
    book.close()
    try:
        reader.read(path, "Лист", source="main", id_column="ID")
    except ProjectError as exc:
        assert exc.diagnostic.code == "FORMULA_ERROR" and exc.diagnostic.cell == "A2"
    else:
        raise AssertionError("Excel error-cell в ID должна блокировать источник")


def test_invalid_dynamic_paths_keep_origin_for_image_and_html(tmp_path: Path) -> None:
    project = ProjectDefinition(schema_version=4, id="game", name="Game", version="0.1.0")
    row = DataRow(2, "2", 1, {"Путь": CellValue("../outside.png", "A2")})
    data = DataSheetSnapshot("main", tmp_path / "data.xlsx", "Лист", "hash", 1, ("Путь",), (row,))
    for node in (
        ImageNode(
            id="art", name="Art", type="image", x_mm=0, y_mm=0,
            width_mm=10, height_mm=10, source="assets/images/{Путь}",
        ),
        HtmlNode(
            id="body", name="Body", type="html", x_mm=0, y_mm=0,
            width_mm=30, height_mm=20, font_family="Arial", font_size_pt=8,
            html='<img src="assets/images/{Путь}" />',
        ),
    ):
        component = ComponentDefinition(
            schema_version=4, id="card", name="Card", size_mm=SizeMM(width=40, height=30),
            data=DataBinding(source="main", sheet="Лист"), elements=(node,),
        )
        try:
            BindingResolver(tmp_path, project).resolve_component(component, data, row)
        except ProjectError as exc:
            diagnostic = exc.diagnostic
            assert diagnostic.code == "PATH_INVALID"
            assert (diagnostic.source, diagnostic.sheet, diagnostic.cell, diagnostic.node_id) == ("main", "Лист", "A2", node.id)
        else:
            raise AssertionError("выход динамического пути должен быть отклонён")


def test_template_unescape_never_changes_inserted_excel_or_variable_data(tmp_path: Path) -> None:
    project = ProjectDefinition(
        schema_version=4, id="game", name="Game", version="0.1.0",
        variables={"literal": r"var \{literal\}"},
    )
    component = ComponentDefinition(
        schema_version=4, id="card", name="Card", size_mm=SizeMM(width=40, height=30),
        data=DataBinding(source="main", sheet="Лист"),
        elements=(HtmlNode(
            id="body", name="Body", type="html", x_mm=0, y_mm=0,
            width_mm=30, height_mm=20, font_family="Arial", font_size_pt=8,
            html=r'<p>template \{literal\}; {Текст}; {{ vars["literal"] }}</p>',
        ),),
    )
    row = DataRow(2, "2", 1, {"Текст": CellValue(r"excel \{literal\}", "A2")})
    data = DataSheetSnapshot("main", tmp_path / "data.xlsx", "Лист", "hash", 1, ("Текст",), (row,))
    resolved = BindingResolver(tmp_path, project).resolve_component(component, data, row)
    html = resolved.component.elements[0].html
    assert html == r"<p>template {literal}; excel \{literal\}; var \{literal\}</p>"


def test_gui_applied_yaml_binding_uses_new_sheet_not_cached_data(tmp_path: Path, qtbot) -> None:
    root = tmp_path / "GUI binding cache"
    shutil.copytree(EXCEL_EXAMPLE, root)
    controller = project_controller()
    window = MainWindow(controller)
    qtbot.addWidget(window)
    window.load_session(project_service().open_session(root))
    window.show()
    tab = window.open_component("event-card")
    assert tab is not None and tab.resolved_component is not None
    assert "Лесной дождь" in repr(tab.resolved_component)

    window._request_mode(tab, "text")
    tab.text_editor.setPlainText(tab.text_editor.toPlainText().replace('sheet: "События"', 'sheet: "Карты"'))
    assert window._apply_draft("event-card")
    window._request_mode(tab, "layout")
    assert tab.resolved_component is not None
    assert "Волк" in repr(tab.resolved_component)
    current = window.session.documents["event-card"].model
    assert controller.preview.data(window.session.snapshot, "event-card", component=current).sheet == "Карты"
    window._clear_tabs(discard=True)
