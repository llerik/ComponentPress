"""Read one immutable XLSX byte sequence with formulas and cached values."""

from __future__ import annotations

from contextlib import contextmanager
from io import BytesIO
import hashlib
import posixpath
import unicodedata
from pathlib import Path
from typing import Iterator
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from componentpress.domain.diagnostics import Diagnostic, ProjectError
from componentpress.domain.values import CellValue, DataRow, DataSheetSnapshot, FormulaState, cell_is_empty, value_to_text


_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
_REL_NS = {"p": "http://schemas.openxmlformats.org/package/2006/relationships"}
_INSTANCE_NAMESPACE = UUID("1f8f7f5c-8b38-5f19-b836-1ce6c606865e")


def normalize_header(value: str) -> str:
    return unicodedata.normalize("NFKC", value.strip()).casefold()


def _formula_xml_root(data: bytes, sheet: str) -> ET.Element | None:
    with BytesIO(data) as stream, ZipFile(stream) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        path = _worksheet_path(workbook, relations, sheet)
        return ET.fromstring(archive.read(path)) if path else None


def _worksheet_path(workbook: ET.Element, relations: ET.Element, sheet: str) -> str | None:
    targets = {item.attrib["Id"]: item.attrib["Target"] for item in relations.findall("p:Relationship", _REL_NS)}
    for item in workbook.findall("m:sheets/m:sheet", _NS):
        if item.attrib.get("name") == sheet:
            target = targets.get(item.attrib.get(f"{{{_NS['r']}}}id", ""))
            if target:
                return _normalize_target(target)
    return None


def _normalize_target(target: str) -> str:
    normalized = target.lstrip("/")
    return normalized if normalized.startswith("xl/") else posixpath.normpath(posixpath.join("xl", normalized))


def _formula_states(root: ET.Element | None) -> dict[str, tuple[FormulaState, str | None]]:
    if root is None:
        return {}
    states = {}
    for cell in root.findall(".//m:c", _NS):
        formula = cell.find("m:f", _NS)
        if formula is None:
            continue
        cached = cell.find("m:v", _NS)
        state = _cache_state(cell, cached)
        states[cell.attrib["r"]] = (state, formula.text)
    return states


def _cache_state(cell: ET.Element, cached: ET.Element | None) -> FormulaState:
    if cell.attrib.get("t") == "e":
        return FormulaState.ERROR
    if cached is None:
        return FormulaState.MISSING
    return FormulaState.EMPTY if cached.text is None else FormulaState.CACHED


def _formula_xml_states(data: bytes, sheet: str) -> dict[str, tuple[FormulaState, str | None]]:
    try:
        return _formula_states(_formula_xml_root(data, sheet))
    except (BadZipFile, KeyError, ET.ParseError) as exc:
        raise ProjectError(Diagnostic("XLSX_INVALID", "не удалось прочитать структуру XLSX", Path("."), sheet=sheet)) from exc


def _close_books(books: list[object]) -> Exception | None:
    error = None
    for book in reversed(books):
        try:
            book.close()
        except Exception as exc:
            error = error or exc
    return error


@contextmanager
def _open_books() -> Iterator[list[object]]:
    books = []
    try:
        yield books
    except BaseException:
        _close_books(books)
        raise
    error = _close_books(books)
    if error is not None:
        raise OSError("не удалось закрыть книгу XLSX") from error


def _open_workbook(stream: BytesIO, books: list[object], **options: bool):
    book = load_workbook(stream, **options)
    books.append(book)
    return book


def _read_bytes(path: Path, source: str) -> bytes:
    try:
        return path.read_bytes()
    except FileNotFoundError as exc:
        raise ProjectError(Diagnostic("DATA_SOURCE_MISSING", "книга XLSX не найдена", path, source=source)) from exc
    except OSError as exc:
        raise ProjectError(Diagnostic("FILE_READ", str(exc), path, source=source)) from exc


def _invalid_xlsx(exc: Exception, path: Path, source: str, sheet: str | None = None) -> ProjectError:
    return ProjectError(Diagnostic("XLSX_INVALID", str(exc), path, source=source, sheet=sheet))


def _cell_value(formula_sheet: Worksheet, value_sheet: Worksheet, states, row: int, column: int) -> CellValue:
    coordinate = f"{get_column_letter(column)}{row}"
    formula_cell = formula_sheet.cell(row, column)
    state, formula = states.get(coordinate, (FormulaState.NONE, None))
    if state is FormulaState.NONE and formula_cell.data_type == "e":
        state = FormulaState.ERROR
    value = value_sheet.cell(row, column).value if state is not FormulaState.NONE else formula_cell.value
    return CellValue(value, coordinate, state, formula)


def _is_empty(cell: CellValue) -> bool:
    if cell.formula_state in (FormulaState.MISSING, FormulaState.ERROR):
        return False
    return cell_is_empty(cell)


def _header_value(sheet: Worksheet, column: int, path: Path, source: str, name: str) -> str | None:
    coordinate = f"{get_column_letter(column)}1"
    cell = sheet.cell(1, column)
    raw = cell.value
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if cell.data_type == "f" or not isinstance(raw, str):
        raise ProjectError(Diagnostic("HEADER_INVALID", "заголовок должен быть непустым текстом", path, source=source, sheet=name, cell=coordinate))
    header = raw.strip()
    if "{" in header or "}" in header:
        raise ProjectError(Diagnostic("HEADER_BRACES", "фигурные скобки в заголовке запрещены", path, source=source, sheet=name, cell=coordinate))
    return header


def _headers(sheet: Worksheet, path: Path, source: str, name: str) -> tuple[list[str], dict[str, str]]:
    populated = (cell.column for cell in sheet._cells.values() if cell.value is not None and cell.row == 1)
    max_column = max(populated, default=1)
    headers, seen = [], {}
    for column in range(1, max_column + 1):
        header = _header_value(sheet, column, path, source, name)
        if header is None:
            break
        normalized = normalize_header(header)
        if normalized in seen:
            coordinate = f"{get_column_letter(column)}1"
            raise ProjectError(Diagnostic("HEADER_DUPLICATE", f"повторный заголовок {header!r}", path, source=source, sheet=name, cell=coordinate))
        seen[normalized] = header
        headers.append(header)
    if not headers:
        raise ProjectError(Diagnostic("HEADER_INVALID", "A1 должен содержать непустой текстовый заголовок", path, source=source, sheet=name, cell="A1"))
    return headers, seen


def _last_data_row(formula_sheet: Worksheet, value_sheet: Worksheet, states, headers: list[str]) -> int:
    populated = (cell for cell in formula_sheet._cells.values() if cell.value is not None and cell.column <= len(headers))
    max_row = max((cell.row for cell in populated), default=1)
    last = 1
    for row in range(2, max_row + 1):
        if all(_is_empty(_cell_value(formula_sheet, value_sheet, states, row, column)) for column in range(1, len(headers) + 1)):
            break
        last = row
    return last


def _reject_merged(sheet: Worksheet, headers: list[str], last_row: int, area: str, path: Path, source: str, name: str) -> None:
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= last_row and merged.max_row >= 1 and merged.min_col <= len(headers) and merged.max_col >= 1:
            raise ProjectError(Diagnostic("MERGED_CELLS_UNSUPPORTED", f"объединённые ячейки пересекают таблицу {area}: {merged}", path, source=source, sheet=name))


def _require_columns(seen: dict[str, str], id_column: str | None, copies_columns: tuple[str, str], path: Path, source: str, sheet: str) -> None:
    required = (*((id_column,) if id_column is not None else ()), *(name for name in copies_columns if name))
    for configured in required:
        if normalize_header(configured) not in seen:
            field = "copies_columns" if configured in copies_columns else "data.id_column"
            raise ProjectError(Diagnostic("COLUMN_UNKNOWN", f"нет столбца {configured!r}", path, field, source=source, sheet=sheet))


def _require_usable(cell: CellValue, path: Path, source: str, sheet: str, column: str) -> None:
    if cell.formula_state is FormulaState.MISSING:
        raise ProjectError(Diagnostic("FORMULA_RESULT_UNAVAILABLE", "нет сохранённого результата формулы", path, column, source=source, sheet=sheet, cell=cell.coordinate))
    if cell.formula_state is FormulaState.ERROR:
        raise ProjectError(Diagnostic("FORMULA_ERROR", f"ошибка Excel: {cell.value}", path, column, source=source, sheet=sheet, cell=cell.coordinate))


def _instance_id(cells: dict[str, CellValue], id_column: str | None, seen: dict[str, str], ids: set[str], row: int, path: Path, source: str, sheet: str) -> str:
    if id_column is None:
        return str(row)
    id_cell = cells[seen[normalize_header(id_column)]]
    _require_usable(id_cell, path, source, sheet, id_column)
    instance_id = value_to_text(id_cell.value)
    if not instance_id:
        raise ProjectError(Diagnostic("INSTANCE_ID_EMPTY", "ID экземпляра не может быть пустым", path, source=source, sheet=sheet, cell=id_cell.coordinate))
    if instance_id in ids:
        raise ProjectError(Diagnostic("INSTANCE_ID_DUPLICATE", f"повторный ID экземпляра {instance_id!r}", path, source=source, sheet=sheet, cell=id_cell.coordinate))
    return instance_id


def _copy_counts(cells: dict[str, CellValue], configured: tuple[str, str], seen: dict[str, str], path: Path, source: str, sheet: str) -> tuple[int, int]:
    counts = []
    for name in configured:
        if not name:
            counts.append(1)
            continue
        actual = seen.get(normalize_header(name))
        if actual is None:
            raise ProjectError(Diagnostic("COLUMN_UNKNOWN", f"нет столбца {name!r}", path, "copies_columns", source=source, sheet=sheet))
        cell = cells[actual]
        _require_usable(cell, path, source, sheet, actual)
        raw = cell.value
        if isinstance(raw, bool) or not isinstance(raw, (int, float)) or int(raw) != raw or raw < 0:
            raise ProjectError(Diagnostic("COPIES_INVALID", "тираж должен быть целым числом от 0", path, "copies_columns", source=source, sheet=sheet, cell=cell.coordinate))
        counts.append(int(raw))
    return counts[0], counts[1]


def _rows(formula_sheet: Worksheet, value_sheet: Worksheet, states, headers: list[str], seen: dict[str, str], id_column: str | None, configured: tuple[str, str], last_row: int, path: Path, source: str, sheet: str) -> tuple[DataRow, ...]:
    rows, ids = [], set()
    for row_number in range(2, last_row + 1):
        cells = {header: _cell_value(formula_sheet, value_sheet, states, row_number, column) for column, header in enumerate(headers, 1)}
        if all(_is_empty(cell) for cell in cells.values()):
            continue
        instance_id = _instance_id(cells, id_column, seen, ids, row_number, path, source, sheet)
        ids.add(instance_id)
        copies, test_copies = _copy_counts(cells, configured, seen, path, source, sheet)
        rows.append(DataRow(row_number, instance_id, copies, cells, test_copies))
    return tuple(rows)


def _snapshot(formula_sheet: Worksheet, value_sheet: Worksheet, states, path: Path, sheet: str, source: str, id_column: str | None, configured: tuple[str, str], version: int, request_id: int, digest: str) -> DataSheetSnapshot:
    headers, seen = _headers(formula_sheet, path, source, sheet)
    last_row = _last_data_row(formula_sheet, value_sheet, states, headers)
    area = f"A1:{get_column_letter(len(headers))}{max(1, last_row)}"
    _reject_merged(formula_sheet, headers, last_row, area, path, source, sheet)
    _require_columns(seen, id_column, configured, path, source, sheet)
    rows = _rows(formula_sheet, value_sheet, states, headers, seen, id_column, configured, last_row, path, source, sheet)
    return DataSheetSnapshot(source, path, sheet, digest, version, tuple(headers), rows, request_id, area)


class XlsxReader:
    """No openpyxl objects escape this adapter and every workbook is closed."""

    def sheet_names(self, path: Path) -> tuple[str, ...]:
        try:
            data = path.read_bytes()
            with BytesIO(data) as stream, _open_books() as books:
                return tuple(_open_workbook(stream, books, read_only=True, data_only=True).sheetnames)
        except (OSError, BadZipFile, ValueError, KeyError) as exc:
            raise _invalid_xlsx(exc, path, "main") from exc

    def read(self, path: Path, sheet: str, *, source: str = "main", id_column: str | None = None, copies_column: str | None = None, version: int = 1, request_id: int = 0, data: bytes | None = None, copies_columns: tuple[str, str] | None = None) -> DataSheetSnapshot:
        if data is None:
            data = _read_bytes(path, source)
        digest = hashlib.sha256(data).hexdigest()
        configured = copies_columns or (copies_column or "", copies_column or "")
        try:
            states = _formula_xml_states(data, sheet)
            with BytesIO(data) as formula_stream, BytesIO(data) as value_stream, _open_books() as books:
                formula_book = _open_workbook(formula_stream, books, read_only=False, data_only=False)
                value_book = _open_workbook(value_stream, books, read_only=False, data_only=True)
                if sheet not in formula_book.sheetnames:
                    raise ProjectError(Diagnostic("SHEET_UNKNOWN", f"лист {sheet!r} не найден", path, source=source, sheet=sheet))
                return _snapshot(formula_book[sheet], value_book[sheet], states, path, sheet, source, id_column, configured, version, request_id, digest)
        except ProjectError:
            raise
        except (BadZipFile, OSError, ValueError, KeyError) as exc:
            raise _invalid_xlsx(exc, path, source, sheet) from exc
